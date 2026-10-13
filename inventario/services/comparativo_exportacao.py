"""Exportação .xlsx do comparativo de contagens.

Usa as mesmas linhas e classificações da tela (`comparativo_contagens`), a
partir de uma única carga de dados, para todas as abas serem consistentes.
"""

import io
import re
from datetime import datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from inventario.services.comparativo_contagens import (
    CELULA_DIVERGENTE,
    CELULA_OK,
    FUSO_ROTULO,
    ORDENS,
    Canal,
    StatusContagem,
    StatusSAP,
    Visao,
    analisar,
    carregar_dados,
    filtrar_linhas,
    montar_linhas_posicao,
    montar_linhas_sku,
    rotulo_ordem,
)

CONTENT_TYPE_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
FORMATO_TEXTO = '@'
FORMATO_NUMERO = '#,##0.###'
FORMATO_DATA = 'dd/mm/yyyy hh:mm'

_FONTE_CABECALHO = Font(bold=True, color='1F2933')
_FUNDO_CABECALHO = PatternFill('solid', fgColor='E7EEF5')
# Mesmas cores da tela.
_FUNDO_OK = PatternFill('solid', fgColor='DCFCE7')
_FUNDO_DIVERGENTE = PatternFill('solid', fgColor='FEE2E2')
_FUNDO_PARCIAL = PatternFill('solid', fgColor='FFEDD5')
_FUNDO_NEUTRO = PatternFill('solid', fgColor='E5E7EB')
_FUNDO_STATUS_CONTAGEM = {
    StatusContagem.OK: _FUNDO_OK,
    StatusContagem.DIVERGENTE: _FUNDO_DIVERGENTE,
    StatusContagem.AGUARDANDO: _FUNDO_NEUTRO,
}
_FUNDO_STATUS_SAP = {
    StatusSAP.OK: _FUNDO_OK,
    StatusSAP.PARCIAL: _FUNDO_PARCIAL,
    StatusSAP.DIVERGENTE: _FUNDO_DIVERGENTE,
    StatusSAP.NEUTRO: _FUNDO_NEUTRO,
}
_FUNDO_CELULA = {CELULA_OK: _FUNDO_OK, CELULA_DIVERGENTE: _FUNDO_DIVERGENTE}

TEXTO = 'texto'
NUMERO = 'numero'
DATA = 'data'


def _data_excel(valor: datetime | None):
    if valor is None:
        return None
    return timezone.localtime(valor).replace(tzinfo=None)


def _escrever(cell, valor, tipo: str, fundo: PatternFill | None = None):
    if tipo == NUMERO:
        if valor is None:
            cell.value = '—'
            cell.data_type = 's'
            cell.alignment = Alignment(horizontal='right')
        else:
            cell.value = valor if isinstance(valor, Decimal) else Decimal(valor)
            cell.number_format = FORMATO_NUMERO
    elif tipo == DATA:
        if valor is None:
            cell.value = 'Não registrado'
            cell.data_type = 's'
        else:
            cell.value = valor
            cell.number_format = FORMATO_DATA
    else:
        # Texto literal: data_type 's' impede que '=', '+', '-' ou '@' no
        # início sejam interpretados como fórmula.
        cell.value = '' if valor is None else str(valor)
        cell.data_type = 's'
        cell.number_format = FORMATO_TEXTO
    if fundo is not None:
        cell.fill = fundo


def _escrever_aba(planilha, colunas, linhas):
    for indice, (titulo, _tipo) in enumerate(colunas, start=1):
        cell = planilha.cell(1, indice, titulo)
        cell.font = _FONTE_CABECALHO
        cell.fill = _FUNDO_CABECALHO
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for numero, valores in enumerate(linhas, start=2):
        for indice, ((_titulo, tipo), valor) in enumerate(zip(colunas, valores), start=1):
            fundo = None
            if isinstance(valor, tuple):
                valor, fundo = valor
            _escrever(planilha.cell(numero, indice), valor, tipo, fundo)
    ultima = get_column_letter(len(colunas))
    planilha.auto_filter.ref = f'A1:{ultima}{max(planilha.max_row, 1)}'
    planilha.freeze_panes = 'A2'
    for indice, (titulo, tipo) in enumerate(colunas, start=1):
        letra = get_column_letter(indice)
        tamanho = len(titulo)
        for cell in planilha[letra][1:]:
            texto = '' if cell.value is None else str(cell.value)
            tamanho = max(tamanho, 16 if tipo == DATA else len(texto))
        planilha.column_dimensions[letra].width = min(max(tamanho + 2, 10), 50)


def _colunas_contagem(ordens, incluir_sap_status: bool, colunas_sap) -> list[tuple[str, str]]:
    colunas = [
        ('Posição', TEXTO), ('Código da posição', TEXTO), ('SKU', TEXTO), ('Descrição', TEXTO),
    ]
    colunas.extend((rotulo, NUMERO) for _chave, rotulo in colunas_sap)
    colunas.extend((rotulo_ordem(o), NUMERO) for o in ordens)
    for ordem in ordens:
        colunas.extend([
            (f'Contagem {ordem} · operadores', TEXTO),
            (f'Contagem {ordem} · primeira data/hora', DATA),
            (f'Contagem {ordem} · última data/hora', DATA),
        ])
    colunas.extend([
        ('Status das contagens', TEXTO),
        ('Dupla coincidente', TEXTO),
        ('Orientação', TEXTO),
        ('Necessita C3', TEXTO),
        ('Diferenças identificadas', TEXTO),
    ])
    if incluir_sap_status:
        colunas.extend([
            ('Status SAP', TEXTO),
            ('Detalhe SAP', TEXTO),
            ('Aviso de distribuição', TEXTO),
        ])
    return colunas


def _valores_contagem(linha, ordens, incluir_sap_status: bool, colunas_sap) -> list:
    valores = [
        ', '.join(linha.nomes_posicao) or '—',
        ', '.join(linha.codigos_posicao) or '—',
        linha.codigo_produto,
        linha.descricao,
    ]
    valores.extend(linha.sap.get(chave) for chave, _rotulo in colunas_sap)
    for ordem in ordens:
        medicao = linha.medicoes[ordem]
        valores.append((medicao.quantidade, _FUNDO_CELULA.get(medicao.celula)))
    for ordem in ordens:
        medicao = linha.medicoes[ordem]
        if medicao.registrada:
            valores.extend([
                medicao.operadores_texto,
                _data_excel(medicao.primeira),
                _data_excel(medicao.ultima),
            ])
        else:
            valores.extend(['Sem lançamento', '', ''])
    valores.extend([
        (linha.contagens.rotulo, _FUNDO_STATUS_CONTAGEM[linha.contagens.status]),
        linha.contagens.dupla,
        linha.contagens.orientacao,
        'Sim' if linha.contagens.necessita_c3 else 'Não',
        ' · '.join(linha.contagens.diferencas),
    ])
    if incluir_sap_status:
        rotulo = linha.status_sap.rotulo + (' (provisório)' if linha.status_sap.provisoria else '')
        valores.extend([
            (rotulo, _FUNDO_STATUS_SAP[linha.status_sap.status]),
            linha.status_sap.detalhe,
            linha.distribuicao_divergente,
        ])
    return valores


def _corrigir_linhas_sem_lancamento(planilha, colunas):
    # Colunas de data vazias por ausência de lançamento ficam em branco
    # (o texto "Sem lançamento" já aparece na coluna de operadores).
    for indice, (_titulo, tipo) in enumerate(colunas, start=1):
        if tipo != DATA:
            continue
        for cell in planilha[get_column_letter(indice)][1:]:
            if cell.value == '':
                cell.value = None


def nome_arquivo(comparativo, canal: str, momento: datetime) -> str:
    base = slugify(comparativo.nome) or f'conferencia-{comparativo.pk}'
    base = re.sub(r'[^A-Za-z0-9_-]', '', base)[:60]
    data = timezone.localtime(momento).strftime('%d-%m-%Y_%H%M')
    return f'Comparativo_{base}_{Canal.ROTULOS[canal]}_{data}.xlsx'


def gerar_pasta(comparativo, *, canal: str, filtros, usuario) -> tuple[Workbook, str]:
    canal = Canal.normalizar(canal)
    dados = carregar_dados(comparativo)
    ordens = [o for o in ORDENS if o in dados.ordens_associadas or o in (1, 2)]
    colunas_sap = Canal.colunas_sap(canal)

    linhas_sku = filtrar_linhas(montar_linhas_sku(dados, canal), filtros)
    filtros_posicao = type(filtros)(
        termo=filtros.termo,
        status_contagem=filtros.status_contagem,
        status_sap='',
        necessita_c3=filtros.necessita_c3,
    )
    linhas_posicao = filtrar_linhas(montar_linhas_posicao(dados), filtros_posicao)
    resultado = analisar(comparativo, canal=canal, visao=Visao.SKU, filtros=filtros, dados=dados)

    pasta = Workbook()
    resumo = pasta.active
    resumo.title = 'Resumo'
    agora = timezone.now()
    nome_usuario = (
        getattr(getattr(usuario, 'perfil_operacional', None), 'nome', '')
        or usuario.get_username()
    )
    linhas_resumo = [
        ('Conferência', comparativo.nome, TEXTO),
        ('Estado operacional', comparativo.get_estado_display(), TEXTO),
    ]
    for ordem in ORDENS:
        contagem = dados.contagens.get(ordem)
        if contagem is None:
            texto = 'Não associada'
        else:
            texto = (
                f'Inventário #{contagem.inventario.pk} · '
                f'{contagem.inventario.get_status_display()}'
                f'{" · provisório" if contagem.provisoria else ""}'
                f' · {contagem.total_lancamentos} lançamento(s)'
            )
        linhas_resumo.append((rotulo_ordem(ordem), texto, TEXTO))
    referencia = comparativo.referencia_sap
    linhas_resumo.extend([
        ('Base SAP', referencia.descricao if referencia else 'Sem referência SAP', TEXTO),
        ('Data da base SAP', _data_excel(referencia.data_importacao_origem) if referencia else None, DATA),
        ('Canal SAP', Canal.DESCRICOES[canal], TEXTO),
        ('Filtros', filtros.descricao(), TEXTO),
        ('Linhas por SKU', len(linhas_sku), NUMERO),
        ('Linhas por posição', len(linhas_posicao), NUMERO),
        ('Contagem OK (todas as linhas)', resultado.resumo.contagem_ok, NUMERO),
        ('Contagens divergentes (todas as linhas)', resultado.resumo.contagem_divergente, NUMERO),
        ('Aguardando contagem (todas as linhas)', resultado.resumo.contagem_aguardando, NUMERO),
        ('Necessita C3 (todas as linhas)', resultado.resumo.necessita_c3, NUMERO),
        ('SAP OK (todas as linhas)', resultado.resumo.sap_ok, NUMERO),
        ('SAP coincide com uma contagem (todas as linhas)', resultado.resumo.sap_parcial, NUMERO),
        ('SAP não coincide (todas as linhas)', resultado.resumo.sap_divergente, NUMERO),
        (
            f'Acuracidade SAP · {Canal.ROTULOS[canal]} (%)',
            resultado.acuracidade.percentual if resultado.acuracidade.disponivel else None,
            NUMERO,
        ),
        ('Acuracidade · critério', 'SKUs com duas contagens iguais; acerto = consenso igual ao saldo SAP. '
         + (resultado.acuracidade.motivo or ''), TEXTO),
        ('SKUs avaliados (com consenso)', resultado.acuracidade.avaliados, NUMERO),
        ('SKUs iguais ao SAP', resultado.acuracidade.acertos, NUMERO),
        ('SKUs diferentes do SAP', resultado.acuracidade.divergentes, NUMERO),
        ('SKUs sem consenso (fora do cálculo)', resultado.acuracidade.sem_consenso, NUMERO),
        ('Exportado por', nome_usuario, TEXTO),
        ('Data/hora da extração', _data_excel(agora), DATA),
        ('Fuso horário', FUSO_ROTULO, TEXTO),
        ('Observação', 'Análise somente leitura: não ajusta estoque nem altera contagens.', TEXTO),
    ])
    resumo.cell(1, 1, 'Campo').font = _FONTE_CABECALHO
    resumo.cell(1, 2, 'Valor').font = _FONTE_CABECALHO
    for numero, (campo, valor, tipo) in enumerate(linhas_resumo, start=2):
        _escrever(resumo.cell(numero, 1), campo, TEXTO)
        _escrever(resumo.cell(numero, 2), valor, tipo)
    resumo.column_dimensions['A'].width = 46
    resumo.column_dimensions['B'].width = 70

    aba_sku = pasta.create_sheet('Por SKU')
    colunas = _colunas_contagem(ordens, True, colunas_sap)
    _escrever_aba(aba_sku, colunas, [
        _valores_contagem(linha, ordens, True, colunas_sap) for linha in linhas_sku
    ])
    _corrigir_linhas_sem_lancamento(aba_sku, colunas)

    aba_posicao = pasta.create_sheet('Por posição')
    colunas = _colunas_contagem(ordens, False, [])
    _escrever_aba(aba_posicao, colunas, [
        _valores_contagem(linha, ordens, False, []) for linha in linhas_posicao
    ])
    _corrigir_linhas_sem_lancamento(aba_posicao, colunas)

    aba_lancamentos = pasta.create_sheet('Lançamentos válidos')
    colunas_lancamento = [
        ('ID', NUMERO),
        ('Contagem', TEXTO),
        ('Inventário', NUMERO),
        ('Posição', TEXTO),
        ('Código da posição', TEXTO),
        ('SKU', TEXTO),
        ('Descrição', TEXTO),
        ('Quantidade', NUMERO),
        ('Embalagem', TEXTO),
        ('Operador original', TEXTO),
        ('Data/hora original', DATA),
        ('Autor da alteração', TEXTO),
        ('Data/hora da alteração', DATA),
        ('Origem', TEXTO),
        ('Situação', TEXTO),
    ]
    linhas_lancamento = []
    for linha in linhas_sku:
        for lancamento in linha.lancamentos:
            linhas_lancamento.append([
                lancamento.id,
                lancamento.rotulo_ordem,
                lancamento.inventario_id,
                lancamento.descricao_posicao or lancamento.codigo_posicao,
                lancamento.codigo_posicao,
                linha.codigo_produto,
                linha.descricao,
                lancamento.quantidade,
                linha.embalagem,
                lancamento.operador,
                _data_excel(lancamento.data_registro),
                lancamento.autor_alteracao,
                _data_excel(lancamento.data_alteracao) if lancamento.data_alteracao else '',
                lancamento.origem,
                lancamento.situacao,
            ])
    _escrever_aba(aba_lancamentos, colunas_lancamento, linhas_lancamento)
    _corrigir_linhas_sem_lancamento(aba_lancamentos, colunas_lancamento)

    aba_acuracidade = pasta.create_sheet('Acuracidade salva')
    colunas_acuracidade = [
        ('Data/hora', DATA),
        ('Canal', TEXTO),
        ('Acuracidade (%)', NUMERO),
        ('SKUs avaliados', NUMERO),
        ('SKUs iguais ao SAP', NUMERO),
        ('SKUs diferentes do SAP', NUMERO),
        ('SKUs sem consenso', NUMERO),
        ('SKUs no comparativo', NUMERO),
        ('Fotografia SAP', TEXTO),
        ('Usuário', TEXTO),
    ]
    registros = comparativo.acuracidades.select_related(
        'usuario', 'usuario__perfil_operacional', 'referencia_sap',
    )
    _escrever_aba(aba_acuracidade, colunas_acuracidade, [
        [
            _data_excel(registro.data_hora),
            Canal.ROTULOS.get(registro.canal, registro.canal),
            registro.percentual,
            registro.skus_avaliados,
            registro.skus_acertos,
            registro.skus_divergentes,
            registro.skus_sem_consenso,
            registro.skus_total,
            f'#{registro.referencia_sap_id}' if registro.referencia_sap_id else '—',
            (
                getattr(getattr(registro.usuario, 'perfil_operacional', None), 'nome', '')
                or (registro.usuario.get_username() if registro.usuario else 'Não registrado')
            ),
        ]
        for registro in registros
    ])

    return pasta, nome_arquivo(comparativo, canal, agora)


def resposta_excel_comparativo(comparativo, *, canal: str, filtros, usuario) -> HttpResponse:
    pasta, nome = gerar_pasta(comparativo, canal=canal, filtros=filtros, usuario=usuario)
    buffer = io.BytesIO()
    pasta.save(buffer)
    resposta = HttpResponse(buffer.getvalue(), content_type=CONTENT_TYPE_XLSX)
    resposta['Content-Disposition'] = f'attachment; filename="{nome}"'
    return resposta
