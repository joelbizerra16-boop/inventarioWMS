import io
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from inventario.services.confronto import CanalConfronto

CONTENT_TYPE_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
FORMATO_TEXTO = '@'
FORMATO_NUMERO = '0.000'
_FONTE_CABECALHO = Font(bold=True, color='1F2933')
_FUNDO_CABECALHO = PatternFill('solid', fgColor='E7EEF5')


def nome_arquivo_exportacao(tela: str, inventario_id: int) -> str:
    data = timezone.localdate().strftime('%d-%m-%Y')
    return f'{tela}_Inventario_{inventario_id}_{data}.xlsx'


def parametros_confronto(params) -> dict:
    """Os mesmos parâmetros GET lidos pela tela de confronto."""
    return {
        'inventario_id': params.get('inventario', '').strip(),
        'filtro_status': params.get('filtro', 'todos'),
        'termo_busca': params.get('q', ''),
        'canal': CanalConfronto.normalizar(params.get('canal', '')),
    }


def colunas_confronto(canal: str) -> list[tuple[str, str]]:
    colunas = [
        ('Código', 'texto'),
        ('Descrição', 'texto'),
        ('Embalagem', 'texto'),
        ('Setor', 'texto'),
        ('Físico', 'numero'),
    ]
    if canal != CanalConfronto.BRIDA:
        colunas.append(('Canal 110 (COSAN)', 'numero'))
    if canal != CanalConfronto.COSAN:
        colunas.append(('Canal 1 (BRIDA)', 'numero'))
    colunas.extend([
        ('Total Contábil', 'numero'),
        ('Diferença', 'numero'),
        ('Status', 'texto'),
    ])
    return colunas


def _texto_ou_traco(valor) -> str:
    texto = '' if valor is None else str(valor).strip()
    return texto or '—'


def _data_contagem(contagem) -> str:
    if not contagem.data_contagem:
        return '—'
    return timezone.localtime(contagem.data_contagem).strftime('%d/%m/%Y %H:%M')


def linhas_contagem(contagens) -> list[list]:
    linhas = []
    for contagem in contagens:
        linhas.append([
            contagem.posicao.codigo,
            contagem.posicao.posicao,
            f'{contagem.produto.codigo_produto} — {contagem.produto.descricao}',
            contagem.quantidade_fisica,
            contagem.usuario_contagem_nome,
            _data_contagem(contagem),
            contagem.origem_contagem_rotulo,
        ])
    return linhas


def linhas_confronto(linhas, canal: str) -> list[list]:
    exportadas = []
    for linha in linhas:
        valores = [
            linha.codigo_produto,
            linha.descricao,
            _texto_ou_traco(linha.embalagem),
            _texto_ou_traco(linha.setor),
            linha.fisico,
        ]
        if canal != CanalConfronto.BRIDA:
            valores.append(linha.canal_110)
        if canal != CanalConfronto.COSAN:
            valores.append(linha.canal_1)
        valores.extend([
            linha.total_contabil,
            linha.diferenca,
            linha.status_label,
        ])
        exportadas.append(valores)
    return exportadas


def _escrever_celula(cell, valor, tipo: str):
    if tipo == 'numero':
        cell.value = valor if isinstance(valor, Decimal) else Decimal(valor or 0)
        cell.number_format = FORMATO_NUMERO
        return
    cell.number_format = FORMATO_TEXTO
    cell.value = '' if valor is None else str(valor)


def _ajustar_larguras(planilha):
    for colunas in planilha.columns:
        letra = get_column_letter(colunas[0].column)
        tamanho = 0
        for cell in colunas:
            texto = '' if cell.value is None else str(cell.value)
            tamanho = max(tamanho, len(texto))
        planilha.column_dimensions[letra].width = min(max(tamanho + 2, 12), 48)


def resposta_planilha(nome_arquivo: str, titulo_aba: str, colunas, linhas) -> HttpResponse:
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = titulo_aba

    for indice, (titulo, _tipo) in enumerate(colunas, start=1):
        cell = planilha.cell(1, indice, titulo)
        cell.font = _FONTE_CABECALHO
        cell.fill = _FUNDO_CABECALHO

    for linha_indice, valores in enumerate(linhas, start=2):
        for coluna_indice, ((_titulo, tipo), valor) in enumerate(zip(colunas, valores), start=1):
            _escrever_celula(planilha.cell(linha_indice, coluna_indice), valor, tipo)

    ultima_coluna = get_column_letter(len(colunas))
    ultima_linha = max(planilha.max_row, 1)
    planilha.auto_filter.ref = f'A1:{ultima_coluna}{ultima_linha}'
    planilha.freeze_panes = 'A2'
    _ajustar_larguras(planilha)

    buffer = io.BytesIO()
    pasta.save(buffer)
    resposta = HttpResponse(buffer.getvalue(), content_type=CONTENT_TYPE_XLSX)
    resposta['Content-Disposition'] = f'attachment; filename="{nome_arquivo}"'
    return resposta


COLUNAS_CONTAGEM = [
    ('Código Alocação', 'texto'),
    ('Alocação', 'texto'),
    ('Produto', 'texto'),
    ('Quantidade', 'numero'),
    ('Usuário', 'texto'),
    ('Data/Hora', 'texto'),
    ('Origem', 'texto'),
]


def resposta_excel_contagem(inventario, contagens) -> HttpResponse:
    return resposta_planilha(
        nome_arquivo_exportacao('Contagem', inventario.pk),
        'Contagem',
        COLUNAS_CONTAGEM,
        linhas_contagem(contagens),
    )


def resposta_excel_confronto(inventario, linhas, canal: str) -> HttpResponse:
    return resposta_planilha(
        nome_arquivo_exportacao('Confronto', inventario.pk),
        'Confronto',
        colunas_confronto(canal),
        linhas_confronto(linhas, canal),
    )
