import io
import re
from datetime import datetime, timedelta, timezone as dt_timezone
from decimal import Decimal

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import Usuario
from accounts.test_utils import criar_usuario_teste
from estoque_sap.models import EstoqueSAP
from inventario.models import Inventario, InventarioItem
from inventario.models_comparativo import (
    ComparativoAcuracidade,
    ComparativoContagem,
    ComparativoContagemEvento,
    ReferenciaSAP,
)
from inventario.services.comparativo_contagens import (
    CELULA_AUSENTE,
    CELULA_DIVERGENTE,
    CELULA_OK,
    Canal,
    ComparativoError,
    FiltrosComparativo,
    StatusContagem,
    StatusSAP,
    Visao,
    analisar,
    capturar_referencia_sap,
    classificar_contagens,
    classificar_sap,
    formatar_quantidade,
    calcular_acuracidade,
    salvar_acuracidade,
    salvar_comparativo,
)
from inventario.services.contagem import excluir_contagem, salvar_contagem
from posicoes.models import Posicao
from produtos.models import Produto

D = Decimal
VERDE, VERMELHO, NEUTRO, LARANJA = 'verde', 'vermelho', 'neutro', 'laranja'
COR_CONTAGEM = {StatusContagem.OK: VERDE, StatusContagem.DIVERGENTE: VERMELHO, StatusContagem.AGUARDANDO: NEUTRO}
COR_SAP = {StatusSAP.OK: VERDE, StatusSAP.PARCIAL: LARANJA, StatusSAP.DIVERGENTE: VERMELHO, StatusSAP.NEUTRO: NEUTRO}


def _q(c1, c2, c3):
    return {1: c1, 2: c2, 3: c3}


class MatrizClassificacaoTestCase(SimpleTestCase):
    """Matriz obrigatória da especificação (SAP selecionado | C1 | C2 | C3)."""

    MATRIZ = [
        (D(100), D(100), D(100), None, VERDE, VERDE),
        (D(100), D(100), D(100), D(90), VERDE, VERDE),
        (D(100), D(100), D(90), D(100), VERDE, VERDE),
        (D(100), D(90), D(100), D(100), VERDE, VERDE),
        (D(100), D(90), D(90), None, VERDE, VERMELHO),
        (D(100), D(90), D(90), D(100), VERDE, LARANJA),
        (D(100), D(100), D(90), None, VERMELHO, LARANJA),
        (D(100), D(90), D(80), None, VERMELHO, VERMELHO),
        (D(100), D(90), D(80), D(70), VERMELHO, VERMELHO),
        (D(100), D(100), None, None, NEUTRO, LARANJA),
        (D(100), None, None, None, NEUTRO, NEUTRO),
        (D(0), D(0), D(0), None, VERDE, VERDE),
    ]

    def test_matriz_obrigatoria(self):
        for sap, c1, c2, c3, esperado_contagem, esperado_sap in self.MATRIZ:
            with self.subTest(sap=sap, c1=c1, c2=c2, c3=c3):
                quantidades = _q(c1, c2, c3)
                contagens = classificar_contagens(quantidades, {1, 2, 3})
                status_sap = classificar_sap(quantidades, sap, contagens)
                self.assertEqual(COR_CONTAGEM[contagens.status], esperado_contagem)
                self.assertEqual(COR_SAP[status_sap.status], esperado_sap)

    def test_c1_igual_c2_dispensa_c3(self):
        resultado = classificar_contagens(_q(D(5), D(5), None), {1, 2, 3})
        self.assertEqual(resultado.dupla, 'C1 = C2')
        self.assertEqual(resultado.orientacao, 'C3 dispensada')
        self.assertFalse(resultado.necessita_c3)
        self.assertEqual(resultado.celulas[3], CELULA_AUSENTE)

    def test_terceira_divergente_preservada_e_destacada(self):
        quantidades = _q(D(100), D(100), D(90))
        resultado = classificar_contagens(quantidades, {1, 2, 3})
        self.assertEqual(resultado.celulas, {1: CELULA_OK, 2: CELULA_OK, 3: CELULA_DIVERGENTE})
        self.assertIn('C3 − C1: −10', resultado.diferencas)
        self.assertEqual(classificar_sap(quantidades, D(100), resultado).status, StatusSAP.OK)

    def test_desempate_identifica_divergente(self):
        c1_c3 = classificar_contagens(_q(D(7), D(8), D(7)), {1, 2, 3})
        self.assertEqual(c1_c3.dupla, 'C1 = C3')
        self.assertEqual(c1_c3.celulas[2], CELULA_DIVERGENTE)
        self.assertIn('C2 divergente', c1_c3.orientacao)
        c2_c3 = classificar_contagens(_q(D(9), D(8), D(8)), {1, 2, 3})
        self.assertEqual(c2_c3.dupla, 'C2 = C3')
        self.assertEqual(c2_c3.celulas[1], CELULA_DIVERGENTE)

    def test_todas_coincidem(self):
        self.assertEqual(classificar_contagens(_q(D(1), D(1), D(1)), {1, 2, 3}).dupla, 'Todas coincidem')

    def test_orientacoes_de_c3(self):
        sem_c3 = classificar_contagens(_q(D(8), D(10), None), {1, 2})
        self.assertEqual(sem_c3.orientacao, 'Solicitar C3')
        self.assertTrue(sem_c3.necessita_c3)
        self.assertIn('C2 − C1: +2', sem_c3.diferencas)
        com_c3 = classificar_contagens(_q(D(8), D(10), None), {1, 2, 3})
        self.assertEqual(com_c3.orientacao, 'Aguardando C3')
        todas = classificar_contagens(_q(D(7), D(8), D(9)), {1, 2, 3})
        self.assertEqual(todas.orientacao, 'Revisão necessária')
        self.assertFalse(todas.necessita_c3)
        self.assertEqual(set(todas.celulas.values()), {CELULA_DIVERGENTE})

    def test_dupla_divergente_do_sap_nao_exige_c3(self):
        quantidades = _q(D(90), D(90), None)
        contagens = classificar_contagens(quantidades, {1, 2})
        self.assertFalse(contagens.necessita_c3)
        self.assertEqual(classificar_sap(quantidades, D(100), contagens).status, StatusSAP.DIVERGENTE)

    def test_resultado_sap_provisorio_com_contagem_pendente(self):
        quantidades = _q(D(100), None, None)
        contagens = classificar_contagens(quantidades, {1, 2})
        self.assertTrue(classificar_sap(quantidades, D(100), contagens).provisoria)

    def test_zero_e_ausencia(self):
        self.assertEqual(formatar_quantidade(None), '—')
        self.assertEqual(formatar_quantidade(D('0')), '0')
        self.assertEqual(formatar_quantidade(D('1234.500')), '1.234,5')
        contagens = classificar_contagens(_q(D(0), None, None), {1, 2})
        self.assertEqual(contagens.status, StatusContagem.AGUARDANDO)
        self.assertEqual(contagens.celulas[1], 'neutra')


class ComparativoBaseTestCase(TestCase):
    def setUp(self):
        self.user, self.operacional = criar_usuario_teste()
        self.client.force_login(self.user)
        self.ana, op_ana = criar_usuario_teste('ana', perfil=Usuario.Perfil.INVENTARIO)
        op_ana.nome = 'Ana'
        op_ana.save(update_fields=['nome'])
        self.bruno, op_bruno = criar_usuario_teste('bruno', perfil=Usuario.Perfil.INVENTARIO)
        op_bruno.nome = 'Bruno'
        op_bruno.save(update_fields=['nome'])
        self.carla, op_carla = criar_usuario_teste('carla', perfil=Usuario.Perfil.INVENTARIO)
        op_carla.nome = 'Carla'
        op_carla.save(update_fields=['nome'])

        self.momento = datetime(2026, 10, 13, 12, 12, tzinfo=dt_timezone.utc)  # 09:12 em Brasília
        self.inv1 = Inventario.objects.create(usuario=self.operacional, status=Inventario.Status.FINALIZADO)
        self.inv2 = Inventario.objects.create(usuario=self.operacional, status=Inventario.Status.EM_ANDAMENTO)
        self.inv3 = Inventario.objects.create(usuario=self.operacional, status=Inventario.Status.ABERTO)
        self.pos_a = Posicao.objects.create(codigo='A-01-01', posicao='Rua A')
        self.pos_b = Posicao.objects.create(codigo='B-02-01', posicao='Rua B')

        self.produto = Produto.objects.create(codigo_produto='00123', descricao='M-SUPER 1000', setor='A')
        self.sap(self.produto, cosan='80', brida='20')

    def sap(self, produto, cosan='0', brida='0'):
        return EstoqueSAP.objects.create(
            produto=produto,
            canal_110=D(cosan),
            canal_1=D(brida),
            total=D(cosan) + D(brida),
            arquivo_origem='sap_teste.xlsx',
        )

    def item(self, inventario, posicao, produto, quantidade, usuario=None, momento=None):
        return InventarioItem.objects.create(
            inventario=inventario,
            posicao=posicao,
            produto=produto,
            quantidade_fisica=D(quantidade),
            usuario_contagem=usuario,
            data_contagem=momento or self.momento,
            usuario_registro=usuario,
            data_registro=momento or self.momento,
        )

    def criar_comparativo(self, c3=True, referencia='nova'):
        return salvar_comparativo(
            comparativo=None,
            nome='Conferência Outubro',
            inventarios={1: self.inv1, 2: self.inv2, 3: self.inv3 if c3 else None},
            referencia_opcao=referencia,
            estado=None,
            usuario=self.user,
        )

    def url(self, comparativo, **params):
        base = reverse('comparativo:acompanhamento', args=[comparativo.pk])
        if params:
            base += '?' + '&'.join(f'{k}={v}' for k, v in params.items())
        return base

    def linha(self, resultado, codigo):
        return next(l for l in resultado.linhas if l.codigo_produto == codigo)


def _cabecalhos_html(conteudo: str) -> list[str]:
    thead = conteudo.split('<thead class="table-light">', 1)[1].split('</thead>', 1)[0]
    cabecalhos = []
    for th in re.findall(r'<th[^>]*>(.*?)</th>', thead, flags=re.S):
        texto = re.sub(r'<div class="cmp-sub">.*?</div>', '', th, flags=re.S)
        texto = re.sub(r'<[^>]+>', '', texto)
        texto = ' '.join(texto.split())
        cabecalhos.append(texto)
    return cabecalhos


class CanaisEColunasTestCase(ComparativoBaseTestCase):
    def setUp(self):
        super().setUp()
        self.item(self.inv1, self.pos_a, self.produto, '100', self.ana)
        self.item(self.inv2, self.pos_a, self.produto, '100', self.bruno)
        self.comparativo = self.criar_comparativo()

    def test_teste_adicional_de_canais(self):
        filtros = FiltrosComparativo()
        cosan = self.linha(analisar(self.comparativo, canal=Canal.COSAN, visao=Visao.SKU, filtros=filtros), '00123')
        brida = self.linha(analisar(self.comparativo, canal=Canal.BRIDA, visao=Visao.SKU, filtros=filtros), '00123')
        ambas = self.linha(analisar(self.comparativo, canal=Canal.AMBAS, visao=Visao.SKU, filtros=filtros), '00123')
        self.assertEqual(cosan.saldo_selecionado, D(80))
        self.assertEqual(cosan.status_sap.status, StatusSAP.DIVERGENTE)
        self.assertEqual(brida.saldo_selecionado, D(20))
        self.assertEqual(brida.status_sap.status, StatusSAP.DIVERGENTE)
        self.assertEqual(ambas.sap, {'cosan': D(80), 'brida': D(20), 'total': D(100)})
        self.assertEqual(ambas.sap['cosan'] + ambas.sap['brida'], ambas.sap['total'])
        self.assertEqual(ambas.status_sap.status, StatusSAP.OK)

    def test_ordem_das_colunas_nos_tres_canais(self):
        esperado_fim = ['C1', 'C2', 'C3', 'Contagem 1', 'Contagem 2', 'Contagem 3',
                        'Status das contagens', 'Status SAP', 'Lançamentos']
        casos = {
            'cosan': ['Posição', 'SKU / descrição', 'COSAN'],
            'brida': ['Posição', 'SKU / descrição', 'BRIDA'],
            'ambas': ['Posição', 'SKU / descrição', 'COSAN', 'BRIDA', 'Total SAP'],
        }
        for canal, inicio in casos.items():
            with self.subTest(canal=canal):
                resposta = self.client.get(self.url(self.comparativo, canal=canal, visao='sku'))
                self.assertEqual(resposta.status_code, 200)
                self.assertEqual(_cabecalhos_html(resposta.content.decode()), inicio + esperado_fim)

    def test_troca_de_aba_nao_mantem_valores_nem_cores(self):
        cosan = self.client.get(self.url(self.comparativo, canal='cosan', visao='sku')).content.decode()
        ambas = self.client.get(self.url(self.comparativo, canal='ambas', visao='sku')).content.decode()
        self.assertNotIn('Total SAP', _cabecalhos_html(cosan))
        corpo_cosan = cosan.split('<tbody>', 1)[1]
        corpo_ambas = ambas.split('<tbody>', 1)[1]
        self.assertIn('cmp-status-divergente">SAP não coincide', corpo_cosan)
        self.assertNotIn('cmp-status-ok">SAP OK', corpo_cosan)
        self.assertIn('cmp-status-ok">SAP OK', corpo_ambas)
        self.assertNotIn('SAP não coincide', corpo_ambas)
        self.assertNotIn('>20<', corpo_cosan)

    def test_visao_por_posicao_nao_tem_sap(self):
        resposta = self.client.get(self.url(self.comparativo, canal='ambas', visao='posicao', status_sap='OK'))
        cabecalhos = _cabecalhos_html(resposta.content.decode())
        self.assertEqual(cabecalhos, [
            'Posição', 'SKU / descrição', 'C1', 'C2', 'C3', 'Contagem 1', 'Contagem 2',
            'Contagem 3', 'Status das contagens', 'Lançamentos',
        ])
        self.assertEqual(resposta.context['filtros'].status_sap, '')
        self.assertIsNone(resposta.context['linhas'][0].status_sap)

    def test_sem_c3_associada_oculta_colunas(self):
        comparativo = self.criar_comparativo(c3=False)
        cabecalhos = _cabecalhos_html(self.client.get(self.url(comparativo, canal='brida')).content.decode())
        self.assertNotIn('C3', cabecalhos)
        self.assertNotIn('Contagem 3', cabecalhos)


class ConsolidacaoTestCase(ComparativoBaseTestCase):
    def test_sku_exclusivo_de_cada_fonte(self):
        so_sap = Produto.objects.create(codigo_produto='0001', descricao='Só SAP', setor='A')
        self.sap(so_sap, cosan='5')
        so_c1 = Produto.objects.create(codigo_produto='0002', descricao='Só C1', setor='A')
        so_c2 = Produto.objects.create(codigo_produto='0003', descricao='Só C2', setor='A')
        self.item(self.inv1, self.pos_a, so_c1, '3', self.ana)
        self.item(self.inv2, self.pos_a, so_c2, '4', self.bruno)
        comparativo = self.criar_comparativo()
        resultado = analisar(comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo())
        codigos = [l.codigo_produto for l in resultado.linhas]
        self.assertEqual(codigos, ['0001', '0002', '0003', '00123'])
        linha_sap = self.linha(resultado, '0001')
        self.assertIsNone(linha_sap.medicoes[1].quantidade)
        self.assertEqual(linha_sap.status_sap.status, StatusSAP.NEUTRO)
        linha_c1 = self.linha(resultado, '0002')
        # Importação completa: sem registro no canal vale zero, sinalizado.
        self.assertEqual(linha_c1.status_sap.status, StatusSAP.DIVERGENTE)
        self.assertIn('Sem registro no SAP selecionado', linha_c1.status_sap.detalhe)
        self.assertEqual(linha_c1.sap['total'], D(0))
        self.assertFalse(linha_c1.sap_presenca['total'])
        self.assertFalse(linha_c1.presente_sap)
        self.assertEqual(self.linha(resultado, '0003').medicoes[2].quantidade, D(4))

    def test_varias_posicoes_sem_multiplicar_saldo_e_distribuicao_divergente(self):
        self.item(self.inv1, self.pos_a, self.produto, '60', self.ana)
        self.item(self.inv1, self.pos_b, self.produto, '40', self.bruno, self.momento.replace(hour=13))
        self.item(self.inv2, self.pos_a, self.produto, '50', self.carla)
        self.item(self.inv2, self.pos_b, self.produto, '50', self.carla)
        comparativo = self.criar_comparativo()
        resultado = analisar(comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo())
        linha = self.linha(resultado, '00123')
        self.assertEqual(linha.medicoes[1].quantidade, D(100))
        self.assertEqual(linha.medicoes[2].quantidade, D(100))
        self.assertEqual(linha.sap['total'], D(100))
        self.assertEqual(linha.codigos_posicao, ['A-01-01', 'B-02-01'])
        self.assertEqual(linha.contagens.status, StatusContagem.OK)
        self.assertEqual(linha.status_sap.status, StatusSAP.OK)
        self.assertIn('distribuição por posição diferente (C1 × C2)', linha.distribuicao_divergente)
        # Vários operadores e intervalo de datas na mesma contagem.
        self.assertEqual(linha.medicoes[1].operadores, ['Ana', 'Bruno'])
        self.assertTrue(linha.medicoes[1].intervalo)
        self.assertEqual(timezone.localtime(linha.medicoes[1].primeira).strftime('%H:%M'), '09:12')
        self.assertEqual(timezone.localtime(linha.medicoes[1].ultima).strftime('%H:%M'), '10:12')

        posicoes = analisar(comparativo, canal='ambas', visao='posicao', filtros=FiltrosComparativo())
        self.assertEqual(len(posicoes.linhas), 2)
        self.assertEqual(posicoes.linhas[0].medicoes[1].quantidade, D(60))
        self.assertEqual(posicoes.linhas[0].contagens.status, StatusContagem.DIVERGENTE)

    def test_posicao_exibe_nome_e_codigo(self):
        self.item(self.inv1, self.pos_a, self.produto, '5', self.ana)
        comparativo = self.criar_comparativo()
        for visao in ('sku', 'posicao'):
            with self.subTest(visao=visao):
                corpo = self.client.get(self.url(comparativo, visao=visao)).content.decode().split('<tbody>', 1)[1]
                self.assertIn('<span class="fw-semibold">Rua A</span>', corpo)
                self.assertIn('Cód. A-01-01', corpo)

    def test_precisao_decimal_e_zeros_a_esquerda(self):
        self.item(self.inv1, self.pos_a, self.produto, '0.125', self.ana)
        self.item(self.inv2, self.pos_a, self.produto, '0.125', self.bruno)
        comparativo = self.criar_comparativo()
        resposta = self.client.get(self.url(comparativo, canal='cosan'))
        self.assertContains(resposta, '00123')
        self.assertContains(resposta, '0,125')
        linha = resposta.context['linhas'][0]
        self.assertEqual(linha.medicoes[1].quantidade, D('0.125'))

    def test_zero_registrado_versus_ausencia_na_tela(self):
        self.item(self.inv1, self.pos_a, self.produto, '0', self.ana)
        comparativo = self.criar_comparativo()
        linha = analisar(comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo()).linhas[0]
        self.assertEqual(linha.medicoes[1].quantidade, D(0))
        self.assertIsNone(linha.medicoes[2].quantidade)
        conteudo = self.client.get(self.url(comparativo)).content.decode()
        self.assertIn('cmp-qtd cmp-neutra" title="C1">0</span>', conteudo)
        self.assertIn('cmp-qtd cmp-ausente" title="C2 · sem lançamento">—</span>', conteudo)

    def test_registro_legado_sem_autoria(self):
        InventarioItem.objects.create(
            inventario=self.inv1, posicao=self.pos_a, produto=self.produto, quantidade_fisica=D(1),
        )
        comparativo = self.criar_comparativo()
        linha = analisar(comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo()).linhas[0]
        self.assertEqual(linha.medicoes[1].operadores_texto, 'Não registrado')
        self.assertIsNone(linha.medicoes[1].primeira)
        self.assertContains(self.client.get(self.url(comparativo)), 'Não registrado')

    def test_correcao_preserva_operador_original_e_exclusao_reflete(self):
        original = salvar_contagem(self.inv2, self.pos_a, self.produto, D(10), usuario_contagem=self.ana)
        corrigido = salvar_contagem(
            self.inv2, self.pos_a, self.produto, D(12),
            item_existente=original, usuario_contagem=self.bruno,
        )
        corrigido.refresh_from_db()
        self.assertEqual(corrigido.usuario_registro_id, self.ana.pk)
        self.assertEqual(corrigido.usuario_alteracao_id, self.bruno.pk)
        self.assertIsNotNone(corrigido.data_alteracao)
        comparativo = self.criar_comparativo()
        linha = analisar(comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo()).linhas[0]
        medicao = linha.medicoes[2]
        self.assertEqual(medicao.quantidade, D(12))
        self.assertEqual(medicao.operadores, ['Ana'])
        self.assertTrue(medicao.alterada)
        self.assertEqual(medicao.lancamentos[0].autor_alteracao, 'Bruno')

        excluir_contagem(corrigido, usuario=self.user)
        linha = analisar(comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo()).linhas[0]
        self.assertIsNone(linha.medicoes[2].quantidade)

    def test_correcao_de_legado_preserva_dados_anteriores(self):
        legado = InventarioItem.objects.create(
            inventario=self.inv1, posicao=self.pos_a, produto=self.produto,
            quantidade_fisica=D(3), usuario_contagem=self.ana, data_contagem=self.momento,
        )
        salvar_contagem(self.inv1, self.pos_a, self.produto, D(4), item_existente=legado, usuario_contagem=self.bruno)
        legado.refresh_from_db()
        self.assertEqual(legado.usuario_registro_id, self.ana.pk)
        self.assertEqual(legado.data_registro, self.momento)
        self.assertEqual(legado.usuario_contagem_id, self.bruno.pk)

    def test_filtros_status_e_necessita_c3(self):
        outro = Produto.objects.create(codigo_produto='0777', descricao='Outro óleo', setor='A')
        self.item(self.inv1, self.pos_a, self.produto, '8', self.ana)
        self.item(self.inv2, self.pos_a, self.produto, '10', self.bruno)
        self.item(self.inv1, self.pos_b, outro, '5', self.ana)
        self.item(self.inv2, self.pos_b, outro, '5', self.bruno)
        comparativo = self.criar_comparativo(c3=False)
        filtros = FiltrosComparativo(necessita_c3=True)
        resultado = analisar(comparativo, canal='ambas', visao='sku', filtros=filtros)
        self.assertEqual([l.codigo_produto for l in resultado.linhas], ['00123'])
        self.assertEqual(resultado.linhas[0].contagens.orientacao, 'Solicitar C3')
        resultado = analisar(comparativo, canal='ambas', visao='posicao', filtros=FiltrosComparativo(termo='b-02'))
        self.assertEqual([l.codigo_produto for l in resultado.linhas], ['0777'])
        resultado = analisar(
            comparativo, canal='ambas', visao='sku',
            filtros=FiltrosComparativo(status_contagem=StatusContagem.OK),
        )
        self.assertEqual([l.codigo_produto for l in resultado.linhas], ['0777'])


class ConfiguracaoTestCase(ComparativoBaseTestCase):
    def test_validacoes_de_associacao(self):
        with self.assertRaisesMessage(ComparativoError, 'já ocupa a ordem C1'):
            salvar_comparativo(
                comparativo=None, nome='X', inventarios={1: self.inv1, 2: self.inv1, 3: None},
                referencia_opcao='', estado=None, usuario=self.user,
            )
        with self.assertRaisesMessage(ComparativoError, 'Associe a C2'):
            salvar_comparativo(
                comparativo=None, nome='X', inventarios={1: self.inv1, 2: None, 3: self.inv3},
                referencia_opcao='', estado=None, usuario=self.user,
            )

    def test_prepara_com_c1_associa_depois_e_audita(self):
        comparativo = salvar_comparativo(
            comparativo=None, nome='Parcial', inventarios={1: self.inv1, 2: None, 3: None},
            referencia_opcao='', estado=None, usuario=self.user,
        )
        self.assertEqual(comparativo.associacoes.count(), 1)
        salvar_comparativo(
            comparativo=comparativo, nome='Parcial', inventarios={1: self.inv2, 2: self.inv1, 3: self.inv3},
            referencia_opcao='', estado=None, usuario=self.user,
        )
        ordens = dict(comparativo.associacoes.values_list('ordem', 'inventario_id'))
        self.assertEqual(ordens, {1: self.inv2.pk, 2: self.inv1.pk, 3: self.inv3.pk})
        acoes = list(ComparativoContagemEvento.objects.filter(comparativo=comparativo).values_list('acao', flat=True))
        self.assertEqual(acoes.count('ASSOCIACAO_ALTERADA'), 4)
        estados = dict(Inventario.objects.values_list('pk', 'status'))
        self.assertEqual(estados[self.inv1.pk], Inventario.Status.FINALIZADO)
        self.assertEqual(estados[self.inv3.pk], Inventario.Status.ABERTO)

    def test_fotografia_sap_nao_muda_com_nova_importacao(self):
        comparativo = self.criar_comparativo()
        self.item(self.inv1, self.pos_a, self.produto, '100', self.ana)
        EstoqueSAP.objects.all().delete()
        self.sap(self.produto, cosan='999', brida='1')
        linha = analisar(comparativo, canal='cosan', visao='sku', filtros=FiltrosComparativo()).linhas[0]
        self.assertEqual(linha.sap['cosan'], D(80))
        nova = capturar_referencia_sap(self.user)
        self.assertNotEqual(nova.pk, comparativo.referencia_sap_id)
        self.assertEqual(capturar_referencia_sap(self.user).pk, nova.pk)
        self.assertEqual(ReferenciaSAP.objects.count(), 2)

    def test_formulario_cria_conferencia(self):
        resposta = self.client.post(reverse('comparativo:criar'), {
            'nome': 'Via tela', 'c1': self.inv1.pk, 'c2': self.inv2.pk, 'c3': '',
            'referencia_sap': 'nova', 'estado': 'ABERTA',
        })
        comparativo = ComparativoContagem.objects.get(nome='Via tela')
        self.assertRedirects(resposta, reverse('comparativo:acompanhamento', args=[comparativo.pk]))
        self.assertEqual(comparativo.usuario_criacao, self.user)
        self.assertIsNotNone(comparativo.referencia_sap)
        lista = self.client.get(reverse('comparativo:lista'))
        self.assertContains(lista, 'Via tela')
        self.assertContains(lista, 'Comparativo de contagens')


class PermissoesTestCase(ComparativoBaseTestCase):
    def test_permissoes(self):
        comparativo = self.criar_comparativo()
        urls = [
            reverse('comparativo:lista'),
            reverse('comparativo:acompanhamento', args=[comparativo.pk]),
            reverse('comparativo:exportar', args=[comparativo.pk]),
            reverse('comparativo:criar'),
        ]
        self.client.logout()
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 302)

        operador, _ = criar_usuario_teste('operador.pocket', perfil=Usuario.Perfil.OPERADOR)
        self.client.force_login(operador)
        for url in urls:
            resposta = self.client.get(url)
            self.assertEqual(resposta.status_code, 302)
            self.assertIn('/pocket/', resposta['Location'])

        consulta, _ = criar_usuario_teste('consulta', perfil=Usuario.Perfil.CONSULTA)
        self.client.force_login(consulta)
        self.assertEqual(self.client.get(urls[1]).status_code, 200)
        self.assertEqual(self.client.get(urls[2]).status_code, 200)
        self.assertRedirects(self.client.get(urls[3]), reverse('comparativo:lista'))
        resposta = self.client.post(urls[3], {'nome': 'Proibida', 'c1': self.inv1.pk})
        self.assertFalse(ComparativoContagem.objects.filter(nome='Proibida').exists())
        self.assertEqual(resposta.status_code, 302)


class AtualizacaoTestCase(ComparativoBaseTestCase):
    def test_parcial_reflete_gravacao_confirmada(self):
        comparativo = self.criar_comparativo()
        url = self.url(comparativo, canal='ambas', visao='posicao', parcial='1')
        vazio = self.client.get(url)
        self.assertNotContains(vazio, 'A-01-01')
        self.assertNotContains(vazio, '<html')
        salvar_contagem(self.inv2, self.pos_a, self.produto, D(7), usuario_contagem=self.carla)
        resposta = self.client.get(url)
        self.assertContains(resposta, 'A-01-01')
        self.assertContains(resposta, 'Carla')
        self.assertEqual(resposta['Cache-Control'], 'no-store')
        self.assertIn('data-url=', self.client.get(self.url(comparativo)).content.decode())


class ExcelTestCase(ComparativoBaseTestCase):
    def setUp(self):
        super().setUp()
        perigoso = Produto.objects.create(codigo_produto='0042', descricao='=1+1', setor='A')
        self.sap(perigoso, cosan='3')
        self.item(self.inv1, self.pos_a, self.produto, '100', self.ana)
        self.item(self.inv2, self.pos_a, self.produto, '100', self.bruno)
        self.item(self.inv3, self.pos_a, self.produto, '90', self.carla)
        self.comparativo = self.criar_comparativo()

    def exportar(self, **params):
        url = reverse('comparativo:exportar', args=[self.comparativo.pk])
        resposta = self.client.get(url, params)
        self.assertEqual(resposta.status_code, 200)
        return resposta, load_workbook(io.BytesIO(resposta.content))

    def test_abas_colunas_e_valores(self):
        resposta, pasta = self.exportar(canal='cosan')
        self.assertEqual(
            pasta.sheetnames,
            ['Resumo', 'Consolidado', 'Por SKU', 'Por posição', 'Lançamentos válidos', 'Acuracidade salva'],
        )
        self.assertIn('Comparativo_conferencia-outubro_COSAN_', resposta['Content-Disposition'])
        aba = pasta['Por SKU']
        cabecalhos = [c.value for c in aba[1]]
        self.assertEqual(cabecalhos[:8], ['Posição', 'Código da posição', 'SKU', 'Descrição', 'COSAN', 'C1', 'C2', 'C3'])
        self.assertNotIn('Total SAP', cabecalhos)
        self.assertNotIn('BRIDA', cabecalhos)
        self.assertEqual(aba.freeze_panes, 'A2')
        self.assertTrue(aba.auto_filter.ref)
        linhas = {row[2].value: row for row in aba.iter_rows(min_row=2)}
        linha = linhas['00123']
        self.assertEqual(linha[0].value, 'Rua A')  # posição de armazenagem
        self.assertEqual(linha[1].value, 'A-01-01')  # código da posição
        self.assertEqual(linha[2].data_type, 's')
        self.assertEqual(D(str(linha[4].value)), D(80))
        self.assertEqual(D(str(linha[5].value)), D(100))
        self.assertEqual(linha[7].fill.fgColor.rgb[-6:], 'FEE2E2')  # C3 divergente
        self.assertEqual(linha[5].fill.fgColor.rgb[-6:], 'DCFCE7')
        self.assertIsInstance(linha[cabecalhos.index('Contagem 1 · primeira data/hora')].value, datetime)
        self.assertEqual(linha[cabecalhos.index('Contagem 1 · primeira data/hora')].value.strftime('%H:%M'), '09:12')
        self.assertEqual(linha[cabecalhos.index('Status SAP')].value, 'SAP não coincide')
        self.assertEqual(linha[cabecalhos.index('Dupla coincidente')].value, 'C1 = C2')
        perigosa = linhas['0042']
        self.assertEqual(perigosa[3].value, '=1+1')
        self.assertEqual(perigosa[3].data_type, 's')
        self.assertEqual(perigosa[5].value, '—')  # ausência não vira zero

        posicao = [c.value for c in pasta['Por posição'][1]]
        self.assertEqual(posicao[:7], ['Posição', 'Código da posição', 'SKU', 'Descrição', 'C1', 'C2', 'C3'])
        self.assertNotIn('Status SAP', posicao)
        self.assertNotIn('COSAN', posicao)
        self.assertEqual(pasta['Lançamentos válidos'].max_row, 4)

    def test_aba_consolidado(self):
        divergente = Produto.objects.create(codigo_produto='0099', descricao='Sem consenso', setor='A')
        self.sap(divergente, cosan='1')
        self.item(self.inv1, self.pos_b, divergente, '1', self.ana)
        self.item(self.inv2, self.pos_b, divergente, '2', self.bruno)
        self.item(self.inv3, self.pos_b, divergente, '3', self.carla)

        _resposta, pasta = self.exportar(canal='cosan')
        aba = pasta['Consolidado']
        self.assertEqual(
            [c.value for c in aba[1]],
            ['SKU', 'Descrição', 'COSAN', 'C1', 'C2', 'C3', 'Consolidado'],
        )
        self.assertEqual(aba.freeze_panes, 'A2')
        self.assertTrue(aba.auto_filter.ref)
        linhas = {row[0].value: row for row in aba.iter_rows(min_row=2)}
        linha = linhas['00123']
        self.assertEqual(linha[1].value, 'M-SUPER 1000')
        self.assertEqual(D(str(linha[2].value)), D(80))
        self.assertEqual(D(str(linha[3].value)), D(100))
        self.assertEqual(D(str(linha[4].value)), D(100))
        self.assertEqual(D(str(linha[5].value)), D(90))
        self.assertEqual(D(str(linha[6].value)), D(100))
        sem_consenso = linhas['0099']
        self.assertEqual(sem_consenso[6].value, 'nao consolidado')
        ausente = linhas['0042']
        self.assertEqual(ausente[1].value, '=1+1')
        self.assertEqual(ausente[1].data_type, 's')
        self.assertEqual(ausente[3].value, '—')
        self.assertEqual(ausente[6].value, 'nao consolidado')
        self.assertEqual(linha[2].number_format, '0')
        decimal = Produto.objects.create(codigo_produto='0083', descricao='Com decimal', setor='A')
        self.sap(decimal, cosan='83.22')
        _resposta, pasta = self.exportar(canal='cosan')
        celula = next(r for r in pasta['Consolidado'].iter_rows(min_row=2) if r[0].value == '0083')[2]
        self.assertEqual(D(str(celula.value)), D('83.22'))
        self.assertNotIn(',', celula.number_format)
        self.assertEqual(celula.number_format, '0"."###')

    def test_ambas_e_correspondencia_com_tela(self):
        _resposta, pasta = self.exportar(canal='ambas')
        aba = pasta['Por SKU']
        cabecalhos = [c.value for c in aba[1]]
        self.assertEqual(cabecalhos[4:7], ['COSAN', 'BRIDA', 'Total SAP'])
        linha = next(r for r in aba.iter_rows(min_row=2) if r[2].value == '00123')
        self.assertEqual(D(str(linha[6].value)), D(100))
        tela = self.client.get(self.url(self.comparativo, canal='ambas'))
        linha_tela = next(l for l in tela.context['linhas'] if l.codigo_produto == '00123')
        self.assertEqual(linha[cabecalhos.index('Status SAP')].value, linha_tela.status_sap.rotulo)
        self.assertEqual(linha[cabecalhos.index('Status das contagens')].value, linha_tela.contagens.rotulo)
        _cabecalhos_tela = _cabecalhos_html(tela.content.decode())
        self.assertEqual(_cabecalhos_tela[2:5], cabecalhos[4:7])

    def test_filtros_aplicados_em_todas_as_linhas(self):
        _resposta, pasta = self.exportar(canal='ambas', q='0042')
        aba = pasta['Por SKU']
        self.assertEqual([r[2].value for r in aba.iter_rows(min_row=2)], ['0042'])
        resumo = {r[0].value: r[1].value for r in pasta['Resumo'].iter_rows(min_row=2)}
        self.assertEqual(resumo['Filtros'], 'Pesquisa: 0042')
        self.assertIn('America/Sao_Paulo', resumo['Fuso horário'])
        self.assertIsInstance(resumo['Data/hora da extração'], datetime)


class AcuracidadeTestCase(ComparativoBaseTestCase):
    """Base: SKUs com consenso; acerto: consenso igual ao SAP do canal."""

    def setUp(self):
        super().setUp()
        # 00123: SAP COSAN 80 / BRIDA 20 / total 100; consenso 100.
        self.item(self.inv1, self.pos_a, self.produto, '100', self.ana)
        self.item(self.inv2, self.pos_a, self.produto, '100', self.bruno)
        # 0200: SAP total 10 (COSAN 10); consenso 7 -> diferente do SAP.
        p2 = Produto.objects.create(codigo_produto='0200', descricao='Dois', setor='A')
        self.sap(p2, cosan='10')
        self.item(self.inv1, self.pos_b, p2, '7', self.ana)
        self.item(self.inv2, self.pos_b, p2, '7', self.bruno)
        # 0300: sem consenso (8 x 9) -> fora do cálculo.
        p3 = Produto.objects.create(codigo_produto='0300', descricao='Três', setor='A')
        self.sap(p3, cosan='8')
        self.item(self.inv1, self.pos_b, p3, '8', self.ana)
        self.item(self.inv2, self.pos_b, p3, '9', self.bruno)
        # 0400: só no SAP, sem contagem -> fora do cálculo.
        p4 = Produto.objects.create(codigo_produto='0400', descricao='Quatro', setor='A')
        self.sap(p4, cosan='5')
        # 0500: contado zero e ausente no SAP -> acerto.
        p5 = Produto.objects.create(codigo_produto='0500', descricao='Cinco', setor='A')
        self.item(self.inv1, self.pos_a, p5, '0', self.ana)
        self.item(self.inv2, self.pos_a, p5, '0', self.bruno)
        self.comparativo = self.criar_comparativo()

    def acuracidade(self, canal):
        return analisar(self.comparativo, canal=canal, visao='sku', filtros=FiltrosComparativo()).acuracidade

    def test_calculo_por_canal(self):
        ambas = self.acuracidade('ambas')
        self.assertEqual((ambas.avaliados, ambas.acertos, ambas.sem_consenso, ambas.total), (3, 2, 2, 5))
        self.assertEqual(ambas.percentual, D('66.67'))
        cosan = self.acuracidade('cosan')  # 100 != 80 -> só o 0500 acerta
        self.assertEqual((cosan.avaliados, cosan.acertos), (3, 1))
        self.assertEqual(cosan.percentual, D('33.33'))

    def test_filtros_nao_alteram_acuracidade(self):
        filtrado = analisar(
            self.comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo(termo='0200'),
        )
        self.assertEqual(len(filtrado.linhas), 1)
        self.assertEqual(filtrado.acuracidade.percentual, D('66.67'))

    def test_sem_referencia_ou_sem_consenso(self):
        sem_ref = self.criar_comparativo(referencia='')
        acc = analisar(sem_ref, canal='ambas', visao='sku', filtros=FiltrosComparativo()).acuracidade
        self.assertFalse(acc.disponivel)
        self.assertIsNone(acc.percentual)
        self.assertEqual(acc.motivo, 'Sem referência SAP')
        vazio = calcular_acuracidade([], 'ambas', True)
        self.assertEqual(vazio.motivo, 'Nenhum SKU com duas contagens iguais')
        with self.assertRaises(ComparativoError):
            salvar_acuracidade(sem_ref, 'ambas', self.user)

    def test_tela_exibe_e_salva(self):
        conteudo = self.client.get(self.url(self.comparativo, canal='ambas', visao='sku')).content.decode()
        self.assertIn('Acuracidade SAP · Ambas', conteudo)
        self.assertIn('<strong>66,67%</strong>', conteudo)
        self.assertIn('Salvar acuracidade', conteudo)
        posicao = self.client.get(self.url(self.comparativo, visao='posicao')).content.decode()
        self.assertNotIn('Acuracidade SAP', posicao)

        url = reverse('comparativo:acuracidade_salvar', args=[self.comparativo.pk])
        resposta = self.client.post(url, {'canal': 'cosan'})
        self.assertRedirects(resposta, self.url(self.comparativo, canal='cosan', visao='sku'))
        registro = ComparativoAcuracidade.objects.get()
        self.assertEqual((registro.canal, registro.percentual, registro.skus_acertos, registro.skus_avaliados),
                         ('cosan', D('33.33'), 1, 3))
        self.assertEqual(registro.usuario, self.user)
        self.assertEqual(registro.referencia_sap_id, self.comparativo.referencia_sap_id)
        self.assertTrue(ComparativoContagemEvento.objects.filter(acao='ACURACIDADE_SALVA').exists())
        lista = self.client.get(reverse('comparativo:lista')).content.decode()
        self.assertIn('33,33%', lista)


    def test_consulta_nao_salva(self):
        consulta, _ = criar_usuario_teste('consulta.acc', perfil=Usuario.Perfil.CONSULTA)
        self.client.force_login(consulta)
        url = reverse('comparativo:acuracidade_salvar', args=[self.comparativo.pk])
        self.client.post(url, {'canal': 'ambas'})
        self.assertFalse(ComparativoAcuracidade.objects.exists())

    def test_excel_resumo_e_historico(self):
        salvar_acuracidade(self.comparativo, 'ambas', self.user)
        resposta = self.client.get(reverse('comparativo:exportar', args=[self.comparativo.pk]), {'canal': 'ambas'})
        pasta = load_workbook(io.BytesIO(resposta.content))
        resumo = {r[0].value: r[1].value for r in pasta['Resumo'].iter_rows(min_row=2)}
        self.assertEqual(D(str(resumo['Acuracidade SAP · Ambas (%)'])), D('66.67'))
        self.assertEqual(resumo['SKUs sem consenso (fora do cálculo)'], 2)
        historico = list(pasta['Acuracidade salva'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(historico), 1)
        self.assertEqual(historico[0][1], 'Ambas')
        self.assertEqual(D(str(historico[0][2])), D('66.67'))
        consolidado = {r[0].value: r for r in pasta['Consolidado'].iter_rows(min_row=2)}
        self.assertEqual(D(str(consolidado['00123'][6].value)), D(100))
        self.assertEqual(consolidado['0300'][6].value, 'nao consolidado')
        self.assertEqual(consolidado['0400'][6].value, 'nao consolidado')
        self.assertEqual(D(str(consolidado['0500'][6].value)), D(0))


class FinalizacaoTestCase(ComparativoBaseTestCase):
    """Salvar = gravar o resultado atual e finalizar a conferência."""

    def setUp(self):
        super().setUp()
        self.item(self.inv1, self.pos_a, self.produto, '100', self.ana)
        self.item(self.inv2, self.pos_a, self.produto, '100', self.bruno)
        self.comparativo = self.criar_comparativo()
        self.url_finalizar = reverse('comparativo:finalizar', args=[self.comparativo.pk])

    def test_salvar_grava_resultado_e_finaliza(self):
        resposta = self.client.post(self.url_finalizar)
        self.assertRedirects(resposta, reverse('comparativo:acompanhamento', args=[self.comparativo.pk]))
        self.comparativo.refresh_from_db()
        self.assertTrue(self.comparativo.finalizado)
        self.assertEqual(self.comparativo.estado, ComparativoContagem.Estado.ENCERRADA)
        self.assertEqual(self.comparativo.usuario_finalizacao, self.user)
        self.assertIsNotNone(self.comparativo.data_finalizacao)
        self.assertEqual(len(self.comparativo.snapshot_resultado['lancamentos']), 2)
        canais = set(ComparativoAcuracidade.objects.values_list('canal', flat=True))
        self.assertEqual(canais, {'cosan', 'brida', 'ambas'})
        self.assertTrue(ComparativoContagemEvento.objects.filter(acao='FINALIZADA').exists())
        # Inventários de origem não mudam de estado.
        self.inv2.refresh_from_db()
        self.assertEqual(self.inv2.status, Inventario.Status.EM_ANDAMENTO)

    def test_resultado_gravado_nao_muda_com_novas_contagens(self):
        antes = analisar(self.comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo())
        self.client.post(self.url_finalizar)
        self.comparativo.refresh_from_db()
        # Depois de salvar: nova contagem, correção e mudança de nome do operador.
        novo = Produto.objects.create(codigo_produto='0999', descricao='Novo', setor='A')
        self.item(self.inv1, self.pos_b, novo, '4', self.ana)
        InventarioItem.objects.filter(inventario=self.inv2, produto=self.produto).update(quantidade_fisica=D('55'))
        depois = analisar(self.comparativo, canal='ambas', visao='sku', filtros=FiltrosComparativo())
        self.assertEqual(
            [(l.codigo_produto, l.medicoes[1].quantidade, l.medicoes[2].quantidade, l.status_sap.status) for l in depois.linhas],
            [(l.codigo_produto, l.medicoes[1].quantidade, l.medicoes[2].quantidade, l.status_sap.status) for l in antes.linhas],
        )
        self.assertEqual(depois.linhas[0].medicoes[1].operadores, ['Ana'])
        self.assertEqual(depois.linhas[0].medicoes[1].primeira, self.momento)
        self.assertEqual(depois.acuracidade.percentual, D('100.00'))
        posicao = analisar(self.comparativo, canal='ambas', visao='posicao', filtros=FiltrosComparativo())
        self.assertEqual(len(posicao.linhas), 1)
        # O Excel também lê o resultado gravado.
        pasta = load_workbook(io.BytesIO(
            self.client.get(reverse('comparativo:exportar', args=[self.comparativo.pk])).content
        ))
        self.assertEqual([r[2].value for r in pasta['Por SKU'].iter_rows(min_row=2)], ['00123'])

    def test_tela_finalizada_sem_atualizacao_nem_configuracao(self):
        self.client.post(self.url_finalizar)
        conteudo = self.client.get(self.url(self.comparativo)).content.decode()
        self.assertIn('Conferência salva e finalizada', conteudo)
        self.assertIn('data-url=""', conteudo)
        self.assertNotIn('Salvar e finalizar', conteudo)
        self.assertNotIn('Salvar acuracidade', conteudo)
        resposta = self.client.get(reverse('comparativo:editar', args=[self.comparativo.pk]))
        self.assertRedirects(resposta, reverse('comparativo:acompanhamento', args=[self.comparativo.pk]))
        with self.assertRaisesMessage(ComparativoError, 'não pode ser alterada'):
            salvar_comparativo(
                comparativo=self.comparativo, nome='X', inventarios={1: self.inv1, 2: None, 3: None},
                referencia_opcao='', estado=None, usuario=self.user,
            )

    def test_salvar_duas_vezes_nao_regrava(self):
        self.client.post(self.url_finalizar)
        self.comparativo.refresh_from_db()
        gravado_em = self.comparativo.data_finalizacao
        resposta = self.client.post(self.url_finalizar, follow=True)
        self.assertContains(resposta, 'já foi salva e finalizada')
        self.comparativo.refresh_from_db()
        self.assertEqual(self.comparativo.data_finalizacao, gravado_em)
        self.assertEqual(ComparativoAcuracidade.objects.count(), 3)

    def test_lista_com_icones_e_confirmacao(self):
        lista = self.client.get(reverse('comparativo:lista')).content.decode()
        self.assertIn('grid-acao grid-acao-abrir', lista)
        self.assertIn('grid-acao grid-acao-editar', lista)
        self.assertIn('grid-acao grid-acao-salvar', lista)
        self.assertIn('Tem certeza que deseja salvar esta contagem?', lista)
        self.client.post(self.url_finalizar, {'voltar': 'lista'})
        lista = self.client.get(reverse('comparativo:lista')).content.decode()
        self.assertIn('Finalizada', lista)
        self.assertIn('grid-acao grid-acao-abrir', lista)
        self.assertIn('grid-acao grid-acao-aviso', lista)
        self.assertIn('Reabrir', lista)
        self.assertNotIn('grid-acao grid-acao-editar', lista)
        self.assertNotIn('grid-acao grid-acao-salvar', lista)

    def test_consulta_nao_finaliza(self):
        consulta, _ = criar_usuario_teste('consulta.fin', perfil=Usuario.Perfil.CONSULTA)
        self.client.force_login(consulta)
        self.client.post(self.url_finalizar)
        self.comparativo.refresh_from_db()
        self.assertFalse(self.comparativo.finalizado)
        lista = self.client.get(reverse('comparativo:lista')).content.decode()
        self.assertNotIn('grid-acao grid-acao-salvar', lista)

    def test_reabrir_descarta_resultado_e_libera_exclusao_do_inventario(self):
        self.client.post(self.url_finalizar)
        self.comparativo.refresh_from_db()
        self.assertTrue(self.comparativo.finalizado)

        bloqueio = self.client.post(
            reverse('inventario:excluir', args=[self.inv2.pk]),
            follow=True,
        )
        self.assertContains(bloqueio, 'Reabra a conferência')
        self.assertTrue(Inventario.objects.filter(pk=self.inv2.pk).exists())

        resposta = self.client.post(
            reverse('comparativo:reabrir', args=[self.comparativo.pk]),
            {'voltar': 'lista'},
            follow=True,
        )
        self.assertContains(resposta, 'reaberta')
        self.comparativo.refresh_from_db()
        self.assertFalse(self.comparativo.finalizado)
        self.assertEqual(self.comparativo.estado, ComparativoContagem.Estado.ABERTA)
        self.assertIsNone(self.comparativo.snapshot_resultado)
        self.assertIsNone(self.comparativo.data_finalizacao)
        self.assertTrue(ComparativoContagemEvento.objects.filter(acao='REABERTA').exists())
        lista = resposta.content.decode()
        self.assertIn('grid-acao grid-acao-editar', lista)
        self.assertNotIn('grid-acao grid-acao-aviso', lista)

        self.client.post(reverse('inventario:excluir', args=[self.inv2.pk]))
        self.assertFalse(Inventario.objects.filter(pk=self.inv2.pk).exists())
        self.inv1.refresh_from_db()
        self.assertEqual(self.inv1.status, Inventario.Status.FINALIZADO)

    def test_consulta_nao_reabre(self):
        self.client.post(self.url_finalizar)
        consulta, _ = criar_usuario_teste('consulta.reabrir', perfil=Usuario.Perfil.CONSULTA)
        self.client.force_login(consulta)
        self.client.post(reverse('comparativo:reabrir', args=[self.comparativo.pk]))
        self.comparativo.refresh_from_db()
        self.assertTrue(self.comparativo.finalizado)
        lista = self.client.get(reverse('comparativo:lista')).content.decode()
        self.assertNotIn('grid-acao grid-acao-aviso', lista)

    def test_nao_reabre_conferencia_fechada_ha_mais_de_7_dias(self):
        self.client.post(self.url_finalizar)
        self.comparativo.refresh_from_db()
        self.comparativo.data_finalizacao = timezone.now() - timedelta(days=7, seconds=1)
        self.comparativo.save(update_fields=['data_finalizacao'])

        resposta = self.client.post(
            reverse('comparativo:reabrir', args=[self.comparativo.pk]),
            {'voltar': 'lista'},
            follow=True,
        )
        self.assertContains(resposta, 'mais de 7 dias')
        self.comparativo.refresh_from_db()
        self.assertTrue(self.comparativo.finalizado)
        self.assertContains(resposta, 'Prazo encerrado')
        self.assertNotIn('grid-acao grid-acao-aviso', resposta.content.decode())

        self.comparativo.data_finalizacao = timezone.now() - timedelta(days=6)
        self.comparativo.save(update_fields=['data_finalizacao'])
        lista = self.client.get(reverse('comparativo:lista')).content.decode()
        self.assertIn('grid-acao grid-acao-aviso', lista)


from estoque_sap.models import ImportacaoEstoqueSAP  # noqa: E402
from estoque_sap.services.importacao_estoque_sap import CAMPOS_CANAIS, importar_dados  # noqa: E402
from inventario.services.comparativo_contagens import (  # noqa: E402
    calcular_cobertura,
    carregar_dados,
    finalizar_comparativo,
    sincronizar_referencia_sap,
)

TODOS_CANAIS = list(CAMPOS_CANAIS)


def _linha_sap(codigo, cosan='0', brida='0', outros=None):
    linha = {campo: '0' for campo in CAMPOS_CANAIS}
    linha.update({'codigo_produto': codigo, 'canal_110': str(cosan), 'canal_1': str(brida)})
    for campo, valor in (outros or {}).items():
        linha[campo] = str(valor)
    linha['total'] = str(sum(D(linha[c]) for c in CAMPOS_CANAIS))
    return linha


class SapPorCanalTestCase(TestCase):
    """Teste obrigatório: COSAN A=10 B=20; BRIDA B=5 C=7; C1 B=25 D=4; C2 B=25 E=6."""

    def setUp(self):
        self.user, self.operacional = criar_usuario_teste()
        self.client.force_login(self.user)
        self.ana, _ = criar_usuario_teste('ana.sap', perfil=Usuario.Perfil.INVENTARIO)
        self.momento = datetime(2026, 10, 14, 12, 0, tzinfo=dt_timezone.utc)
        self.inv1 = Inventario.objects.create(usuario=self.operacional, status=Inventario.Status.EM_ANDAMENTO)
        self.inv2 = Inventario.objects.create(usuario=self.operacional, status=Inventario.Status.EM_ANDAMENTO)
        self.pos = Posicao.objects.create(codigo='11', posicao='01.03.101')
        self.pos2 = Posicao.objects.create(codigo='12', posicao='01.03.102')
        self.p = {
            codigo: Produto.objects.create(codigo_produto=codigo, descricao=f'Produto {codigo}', setor='A')
            for codigo in ('A', 'B', 'C', 'D', 'E')
        }
        importar_dados(
            [_linha_sap('A', cosan=10), _linha_sap('B', cosan=20, brida=5), _linha_sap('C', brida=7)],
            arquivo_origem='sap_completo.xlsx',
            canais_carregados=TODOS_CANAIS,
        )
        self.item(self.inv1, 'B', '25')
        self.item(self.inv1, 'D', '4')
        self.item(self.inv2, 'B', '25')
        self.item(self.inv2, 'E', '6')
        self.comparativo = salvar_comparativo(
            comparativo=None, nome='Canais', inventarios={1: self.inv1, 2: self.inv2, 3: None},
            referencia_opcao='nova', estado=None, usuario=self.user,
        )

    def item(self, inventario, codigo, quantidade, posicao=None):
        return InventarioItem.objects.create(
            inventario=inventario, posicao=posicao or self.pos, produto=self.p[codigo],
            quantidade_fisica=D(quantidade), usuario_contagem=self.ana, data_contagem=self.momento,
            usuario_registro=self.ana, data_registro=self.momento,
        )

    def resultado(self, canal):
        return analisar(self.comparativo, canal=canal, visao='sku', filtros=FiltrosComparativo())

    def linhas(self, canal):
        return {l.codigo_produto: l for l in self.resultado(canal).linhas}

    def url(self, **params):
        return reverse('comparativo:acompanhamento', args=[self.comparativo.pk]) + '?' + '&'.join(
            f'{k}={v}' for k, v in params.items()
        )

    def test_lote_registra_canais_e_referencia_herda(self):
        importacao = ImportacaoEstoqueSAP.objects.get()
        self.assertIn('canal_110', importacao.canais_carregados)
        self.assertIn('canal_1', importacao.canais_carregados)
        self.assertEqual(self.comparativo.referencia_sap.canais_carregados, importacao.canais_carregados)

    def test_aba_ambas(self):
        linhas = self.linhas('ambas')
        self.assertEqual(sorted(linhas), ['A', 'B', 'C', 'D', 'E'])
        a, b, c = linhas['A'], linhas['B'], linhas['C']
        self.assertEqual((a.sap['cosan'], a.sap['brida'], a.sap['total']), (D(10), D(0), D(10)))
        self.assertEqual((a.sap_presenca['cosan'], a.sap_presenca['brida']), (True, False))
        self.assertEqual((b.sap['cosan'], b.sap['brida'], b.sap['total']), (D(20), D(5), D(25)))
        self.assertEqual((c.sap['cosan'], c.sap['brida'], c.sap['total']), (D(0), D(7), D(7)))
        self.assertEqual(b.contagens.status, StatusContagem.OK)
        self.assertEqual(b.status_sap.status, StatusSAP.OK)
        for codigo in ('D', 'E'):
            self.assertFalse(linhas[codigo].presente_sap)
            self.assertTrue(linhas[codigo].presente_fisico)
            self.assertIn('Sem registro no SAP selecionado', linhas[codigo].status_sap.detalhe)
        for codigo in ('A', 'C'):
            self.assertIsNone(linhas[codigo].medicoes[1].quantidade)
            self.assertEqual(linhas[codigo].contagens.status, StatusContagem.AGUARDANDO)
            self.assertEqual(linhas[codigo].status_sap.status, StatusSAP.NEUTRO)  # não contado não fica vermelho
        cob = self.resultado('ambas').cobertura
        self.assertEqual(
            (cob.skus_cosan, cob.skus_brida, cob.skus_sap_selecionado, cob.skus_fisicos,
             cob.intersecao, cob.somente_sap, cob.somente_fisico, cob.total),
            (2, 2, 3, 3, 1, 2, 2, 5),
        )

    def test_aba_cosan(self):
        linhas = self.linhas('cosan')
        self.assertEqual(sorted(linhas), ['A', 'B', 'D', 'E'])
        b = linhas['B']
        self.assertEqual(b.saldo_selecionado, D(20))
        self.assertEqual(b.contagens.status, StatusContagem.OK)
        self.assertEqual(b.status_sap.status, StatusSAP.DIVERGENTE)
        cob = self.resultado('cosan').cobertura
        self.assertEqual((cob.skus_sap_selecionado, cob.total, cob.intersecao), (2, 4, 1))
        cabecalhos = _cabecalhos_html(self.client.get(self.url(canal='cosan', visao='sku')).content.decode())
        self.assertEqual(cabecalhos[:4], ['Posição', 'SKU / descrição', 'COSAN', 'C1'])
        self.assertNotIn('BRIDA', cabecalhos)
        self.assertNotIn('Total SAP', cabecalhos)

    def test_aba_brida(self):
        linhas = self.linhas('brida')
        self.assertEqual(sorted(linhas), ['B', 'C', 'D', 'E'])
        b = linhas['B']
        self.assertEqual(b.saldo_selecionado, D(5))
        self.assertEqual(b.contagens.status, StatusContagem.OK)
        self.assertEqual(b.status_sap.status, StatusSAP.DIVERGENTE)
        cabecalhos = _cabecalhos_html(self.client.get(self.url(canal='brida', visao='sku')).content.decode())
        self.assertEqual(cabecalhos[:4], ['Posição', 'SKU / descrição', 'BRIDA', 'C1'])
        self.assertNotIn('COSAN', cabecalhos)
        self.assertNotIn('Total SAP', cabecalhos)

    def test_tela_mostra_sem_registro_e_sem_posicao(self):
        html = self.client.get(self.url(canal='ambas', visao='sku')).content.decode()
        corpo = html.split('<tbody>', 1)[1]
        self.assertIn('sem registro', corpo)
        self.assertIn('Sem posição contada', corpo)
        self.assertIn('SKUs distintos', html)
        self.assertIn('Total do comparativo <strong>5</strong>', html)

    def test_conferencia_sem_contagens_exibe_skus_sap(self):
        vazio = Inventario.objects.create(usuario=self.operacional, status=Inventario.Status.ABERTO)
        comp = salvar_comparativo(
            comparativo=None, nome='Vazia', inventarios={1: vazio, 2: None, 3: None},
            referencia_opcao='nova', estado=None, usuario=self.user,
        )
        for canal, esperado in (('cosan', ['A', 'B']), ('brida', ['B', 'C']), ('ambas', ['A', 'B', 'C'])):
            linhas = analisar(comp, canal=canal, visao='sku', filtros=FiltrosComparativo()).linhas
            self.assertEqual([l.codigo_produto for l in linhas], esperado)

    def test_nova_importacao_entra_auditada_e_preserva_anterior(self):
        anterior = self.comparativo.referencia_sap_id
        novo = Produto.objects.create(codigo_produto='F', descricao='Novo', setor='A')
        importar_dados(
            [_linha_sap('A', cosan=10), _linha_sap('B', cosan=30, brida=5), _linha_sap('C', brida=7),
             _linha_sap('F', brida=3)],
            arquivo_origem='sap_novo.xlsx', canais_carregados=TODOS_CANAIS,
        )
        self.client.get(self.url(canal='ambas', visao='sku'))  # leitura sincroniza
        self.comparativo.refresh_from_db()
        self.assertNotEqual(self.comparativo.referencia_sap_id, anterior)
        self.assertTrue(ReferenciaSAP.objects.filter(pk=anterior).exists())
        evento = ComparativoContagemEvento.objects.get(acao='REFERENCIA_SAP_ATUALIZADA')
        self.assertEqual((evento.dados['antes'], evento.dados['depois']), (anterior, self.comparativo.referencia_sap_id))
        linhas = self.linhas('ambas')
        self.assertIn('F', linhas)
        self.assertEqual(linhas['B'].sap['total'], D(35))  # COSAN 30 + BRIDA 5 recalculado
        self.assertEqual(linhas['B'].status_sap.status, StatusSAP.DIVERGENTE)
        self.assertEqual(novo.pk, linhas['F'].produto_id)

    def test_importacao_parcial_nao_substitui_referencia_completa(self):
        anterior = self.comparativo.referencia_sap_id
        importar_dados(
            [_linha_sap('A', cosan=99)], arquivo_origem='so_cosan.xlsx',
            canais_carregados=['canal_110'],
        )
        alterou, aviso = sincronizar_referencia_sap(self.comparativo, self.user)
        self.assertFalse(alterou)
        self.assertIn('não trouxe BRIDA', aviso)
        self.comparativo.refresh_from_db()
        self.assertEqual(self.comparativo.referencia_sap_id, anterior)
        html = self.client.get(self.url(canal='ambas', visao='sku')).content.decode()
        self.assertIn('referência anterior completa foi mantida', html)
        self.assertEqual(self.linhas('ambas')['A'].sap['cosan'], D(10))

    def test_importacao_parcial_sem_referencia_marca_canal_nao_carregado(self):
        ComparativoContagem.objects.filter(pk=self.comparativo.pk).update(referencia_sap=None)
        self.comparativo.refresh_from_db()
        importar_dados(
            [_linha_sap('A', cosan=10), _linha_sap('B', cosan=20)], arquivo_origem='so_cosan.xlsx',
            canais_carregados=['canal_110'],
        )
        alterou, _aviso = sincronizar_referencia_sap(self.comparativo, self.user)
        self.assertTrue(alterou)
        brida = self.resultado('brida')
        b = {l.codigo_produto: l for l in brida.linhas}['B']
        self.assertIsNone(b.sap['brida'])  # não carregado nunca vira zero
        self.assertIsNone(b.sap['total'])
        self.assertEqual(b.status_sap.status, StatusSAP.NEUTRO)
        self.assertEqual(b.status_sap.rotulo, 'Canal BRIDA não carregado na importação SAP')
        self.assertFalse(brida.acuracidade.disponivel)
        ambas = {l.codigo_produto: l for l in self.resultado('ambas').linhas}
        self.assertEqual(ambas['B'].status_sap.rotulo, 'Base SAP parcial: BRIDA não carregado na importação')
        cosan = {l.codigo_produto: l for l in self.resultado('cosan').linhas}
        self.assertEqual(cosan['B'].status_sap.status, StatusSAP.DIVERGENTE)  # COSAN válido

    def test_varias_posicoes_e_registros_sem_duplicar_saldo(self):
        self.item(self.inv1, 'A', '4', posicao=self.pos)
        self.item(self.inv1, 'A', '6', posicao=self.pos2)
        self.item(self.inv2, 'A', '10', posicao=self.pos2)
        a = self.linhas('ambas')['A']
        self.assertEqual(a.sap['cosan'], D(10))
        self.assertEqual(a.sap['total'], D(10))
        self.assertEqual(a.medicoes[1].quantidade, D(10))
        self.assertEqual(a.status_sap.status, StatusSAP.OK)
        self.assertEqual(len(a.codigos_posicao), 2)
        dados = carregar_dados(self.comparativo)
        self.assertEqual(sum(1 for l in self.resultado('ambas').linhas if l.codigo_produto == 'A'), 1)
        self.assertEqual(calcular_cobertura(dados, 'ambas').intersecao, 2)

    def test_novo_sku_fisico_e_versao_do_polling(self):
        url = self.url(canal='ambas', visao='sku', parcial='1')
        primeira = self.client.get(url)
        versao = primeira['X-Comparativo-Versao']
        sem_mudanca = self.client.get(url + f'&versao={versao}')
        self.assertEqual(sem_mudanca.status_code, 204)
        forcado = self.client.get(url + f'&versao={versao}&forcar=1')
        self.assertEqual(forcado.status_code, 200)
        novo = Produto.objects.create(codigo_produto='G', descricao='Contado agora', setor='A')
        self.p['G'] = novo
        self.item(self.inv2, 'G', '2')
        depois = self.client.get(url + f'&versao={versao}')
        self.assertEqual(depois.status_code, 200)
        self.assertNotEqual(depois['X-Comparativo-Versao'], versao)
        self.assertIn('>G<', depois.content.decode())

    def test_inativo_e_outros_canais_excluidos_com_contagem(self):
        inativo = Produto.objects.create(codigo_produto='H', descricao='Inativo', setor='A', ativo=False)
        outro = Produto.objects.create(codigo_produto='J', descricao='Outro canal', setor='A')
        sem_desc = Produto.objects.create(codigo_produto='0042', descricao='', setor='A')
        importar_dados(
            [_linha_sap('A', cosan=10), _linha_sap('B', cosan=20, brida=5), _linha_sap('C', brida=7),
             _linha_sap('H', cosan=3), _linha_sap('J', outros={'canal_66': 9}), _linha_sap('0042', brida=1)],
            arquivo_origem='sap_completo2.xlsx', canais_carregados=TODOS_CANAIS,
        )
        sincronizar_referencia_sap(self.comparativo, self.user)
        resultado = self.resultado('ambas')
        codigos = [l.codigo_produto for l in resultado.linhas]
        self.assertNotIn('H', codigos)
        self.assertNotIn('J', codigos)
        self.assertEqual(resultado.cobertura.sap_inativos_excluidos, 1)
        self.assertEqual(resultado.cobertura.sap_fora_canais, 1)
        linha = {l.codigo_produto: l for l in resultado.linhas}['0042']
        self.assertEqual(linha.descricao, 'Descrição não cadastrada')
        self.assertEqual(inativo.pk, Produto.objects.get(codigo_produto='H').pk)
        self.assertEqual(outro.pk, Produto.objects.get(codigo_produto='J').pk)
        self.assertEqual(sem_desc.codigo_produto, '0042')

    def test_finalizada_preserva_referencia_historica(self):
        finalizar_comparativo(self.comparativo, self.user)
        self.comparativo.refresh_from_db()
        anterior = self.comparativo.referencia_sap_id
        importar_dados(
            [_linha_sap('A', cosan=1)], arquivo_origem='posterior.xlsx', canais_carregados=TODOS_CANAIS,
        )
        self.client.get(self.url(canal='ambas', visao='sku'))
        self.comparativo.refresh_from_db()
        self.assertEqual(self.comparativo.referencia_sap_id, anterior)
        linhas = self.linhas('ambas')
        self.assertEqual(sorted(linhas), ['A', 'B', 'C', 'D', 'E'])
        self.assertEqual(linhas['A'].sap['cosan'], D(10))
        cob = self.resultado('ambas').cobertura
        self.assertEqual((cob.skus_sap_selecionado, cob.total), (3, 5))

    def test_excel_igual_a_tela(self):
        for canal, colunas_sap in (('cosan', ['COSAN']), ('brida', ['BRIDA']), ('ambas', ['COSAN', 'BRIDA', 'Total SAP'])):
            with self.subTest(canal=canal):
                resposta = self.client.get(
                    reverse('comparativo:exportar', args=[self.comparativo.pk]), {'canal': canal},
                )
                pasta = load_workbook(io.BytesIO(resposta.content))
                aba = pasta['Por SKU']
                cab = [c.value for c in aba[1]]
                self.assertEqual(cab[4:4 + len(colunas_sap)], colunas_sap)
                excel = sorted(r[2].value for r in aba.iter_rows(min_row=2))
                tela = sorted(l.codigo_produto for l in self.resultado(canal).linhas)
                self.assertEqual(excel, tela)
                resumo = {r[0].value: r[1].value for r in pasta['Resumo'].iter_rows(min_row=2)}
                cob = self.resultado(canal).cobertura
                self.assertEqual(resumo['Total de SKUs do comparativo'], cob.total)
                self.assertEqual(resumo['SKUs SAP únicos da seleção'], cob.skus_sap_selecionado)
                self.assertEqual(resumo['Canal BRIDA (1)'], 'carregado')
        pasta = load_workbook(io.BytesIO(self.client.get(
            reverse('comparativo:exportar', args=[self.comparativo.pk]), {'canal': 'ambas'},
        ).content))
        linhas = {r[2].value: r for r in pasta['Por SKU'].iter_rows(min_row=2)}
        self.assertEqual(linhas['A'][5].value, 'sem registro')  # BRIDA de A
        self.assertEqual(D(str(linhas['A'][6].value)), D(10))  # Total SAP de A
        self.assertEqual(linhas['D'][6].value, 'sem registro')

    def test_consulta_visualiza_e_sincroniza(self):
        consulta, _ = criar_usuario_teste('consulta.sap', perfil=Usuario.Perfil.CONSULTA)
        self.client.force_login(consulta)
        self.assertEqual(self.client.get(self.url(canal='ambas', visao='sku')).status_code, 200)
        self.assertEqual(
            self.client.get(reverse('comparativo:exportar', args=[self.comparativo.pk])).status_code, 200,
        )
