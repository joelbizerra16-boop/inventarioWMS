import io
from datetime import datetime
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import Usuario
from accounts.test_utils import ClienteAutenticadoMixin, criar_usuario_teste
from estoque_sap.models import EstoqueSAP
from inventario.models import Inventario, InventarioItem
from inventario.services.exportacao_excel import FORMATO_NUMERO, FORMATO_TEXTO
from posicoes.models import Posicao
from produtos.models import Produto


def _planilha(response):
    pasta = load_workbook(io.BytesIO(response.content))
    return pasta.active


def _cabecalhos(planilha):
    return [cell.value for cell in planilha[1]]


def _linhas(planilha):
    return [
        tuple(cell.value for cell in linha)
        for linha in planilha.iter_rows(min_row=2)
    ]


def _decimal(valor) -> Decimal:
    return Decimal(str(valor))


class ExportacaoExcelTestCase(ClienteAutenticadoMixin, TestCase):
    def setUp(self):
        self.autenticar_cliente()
        self.usuario = Usuario.objects.get(login='admin.teste')
        self.auth_user = self.usuario.user
        self.momento = timezone.make_aware(datetime(2026, 10, 13, 16, 5))

        self.produto_zeros = Produto.objects.create(
            codigo_produto='00123',
            descricao='Parafuso',
            setor='A',
            embalagem='',
        )
        self.produto_outro = Produto.objects.create(
            codigo_produto='00999',
            descricao='Porca',
            setor='',
            embalagem='CX',
        )
        self.produto_so_sap = Produto.objects.create(
            codigo_produto='00055',
            descricao='Arruela',
            setor='B',
            embalagem='UN',
        )
        self.posicao_a = Posicao.objects.create(codigo='0007', posicao='RUA-01')
        self.posicao_b = Posicao.objects.create(codigo='0008', posicao='RUA-02')
        self.posicao_c = Posicao.objects.create(codigo='0009', posicao='RUA-03')
        self.inventario = Inventario.objects.create(
            usuario=self.usuario,
            status=Inventario.Status.EM_ANDAMENTO,
        )
        self.item_a = InventarioItem.objects.create(
            inventario=self.inventario,
            posicao=self.posicao_a,
            produto=self.produto_zeros,
            quantidade_fisica=Decimal('10.000'),
            usuario_contagem=self.auth_user,
            data_contagem=self.momento,
            origem_contagem=InventarioItem.OrigemContagem.WEB,
        )
        self.item_b = InventarioItem.objects.create(
            inventario=self.inventario,
            posicao=self.posicao_b,
            produto=self.produto_zeros,
            quantidade_fisica=Decimal('2.500'),
            usuario_contagem=None,
            data_contagem=None,
            origem_contagem='',
        )
        self.item_c = InventarioItem.objects.create(
            inventario=self.inventario,
            posicao=self.posicao_c,
            produto=self.produto_outro,
            quantidade_fisica=Decimal('3.000'),
        )
        EstoqueSAP.objects.create(
            produto=self.produto_zeros,
            canal_110=Decimal('12.500'),
            canal_1=Decimal('10.000'),
            total=Decimal('99.000'),
            arquivo_origem='teste.xlsx',
        )
        EstoqueSAP.objects.create(
            produto=self.produto_outro,
            canal_110=Decimal('0.000'),
            canal_1=Decimal('0.000'),
            total=Decimal('0.000'),
            arquivo_origem='teste.xlsx',
        )
        EstoqueSAP.objects.create(
            produto=self.produto_so_sap,
            canal_110=Decimal('4.000'),
            canal_1=Decimal('1.000'),
            total=Decimal('5.000'),
            arquivo_origem='teste.xlsx',
        )

    def _exportar_contagem(self, **params):
        return self.client.get(
            reverse('inventario:contagem_exportar', args=[self.inventario.pk]),
            params,
        )

    def _exportar_confronto(self, **params):
        params.setdefault('inventario', self.inventario.pk)
        return self.client.get(reverse('confronto_exportar'), params)

    def test_contagem_exporta_filtro_completo_com_tipos_e_ordem_da_tela(self):
        antes_itens = list(
            InventarioItem.objects.values_list('pk', 'quantidade_fisica')
        )
        params = {'q': '00123', 'page': '2'}
        tela = self.client.get(
            reverse('inventario:contagem_lista', args=[self.inventario.pk]),
            params,
        )
        resposta = self._exportar_contagem(**params)

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta['Content-Type'], (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        ))
        self.assertIn(
            f'Contagem_Inventario_{self.inventario.pk}_',
            resposta['Content-Disposition'],
        )
        self.assertIn(
            timezone.localdate().strftime('%d-%m-%Y'),
            resposta['Content-Disposition'],
        )
        self.assertContains(tela, 'Exportar Excel')
        self.assertContains(tela, 'contagem/exportar/')
        self.assertContains(tela, 'q=00123')

        planilha = _planilha(resposta)
        self.assertEqual(_cabecalhos(planilha), [
            'Código Alocação',
            'Alocação',
            'Produto',
            'Quantidade',
            'Usuário',
            'Data/Hora',
            'Origem',
        ])
        self.assertNotIn('Ações', _cabecalhos(planilha))
        self.assertEqual(planilha.freeze_panes, 'A2')
        self.assertEqual(planilha.auto_filter.ref, 'A1:G3')
        self.assertGreaterEqual(planilha.column_dimensions['A'].width, 12)
        self.assertGreaterEqual(planilha.column_dimensions['C'].width, 12)

        contagens = list(tela.context['contagens'])
        self.assertEqual(len(contagens), 2)
        self.assertEqual(len(_linhas(planilha)), 2)
        self.assertEqual(
            {planilha.cell(indice, 1).value for indice in range(2, 4)},
            {'0007', '0008'},
        )

        for indice, contagem in enumerate(contagens, start=2):
            codigo = planilha.cell(indice, 1)
            produto = planilha.cell(indice, 3)
            quantidade = planilha.cell(indice, 4)
            self.assertEqual(codigo.value, contagem.posicao.codigo)
            self.assertEqual(codigo.number_format, FORMATO_TEXTO)
            self.assertEqual(codigo.data_type, 's')
            self.assertEqual(planilha.cell(indice, 2).value, contagem.posicao.posicao)
            self.assertEqual(
                produto.value,
                f'{contagem.produto.codigo_produto} — {contagem.produto.descricao}',
            )
            self.assertEqual(produto.number_format, FORMATO_TEXTO)
            self.assertTrue(str(produto.value).startswith('00123'))
            self.assertEqual(quantidade.data_type, 'n')
            self.assertEqual(quantidade.number_format, FORMATO_NUMERO)
            self.assertEqual(_decimal(quantidade.value), contagem.quantidade_fisica)
            self.assertEqual(planilha.cell(indice, 5).value, contagem.usuario_contagem_nome)
            if contagem.data_contagem:
                self.assertEqual(
                    planilha.cell(indice, 6).value,
                    timezone.localtime(contagem.data_contagem).strftime('%d/%m/%Y %H:%M'),
                )
            else:
                self.assertEqual(planilha.cell(indice, 6).value, '—')
            self.assertEqual(
                planilha.cell(indice, 7).value,
                contagem.origem_contagem_rotulo,
            )

        self.assertEqual(
            list(InventarioItem.objects.values_list('pk', 'quantidade_fisica')),
            antes_itens,
        )

    def test_confronto_respeita_canal_filtro_e_pesquisa(self):
        antes_sap = list(EstoqueSAP.objects.values_list('pk', 'canal_110', 'canal_1', 'total'))
        antes_itens = InventarioItem.objects.count()

        for canal, coluna_presente, coluna_ausente in (
            ('110', 'Canal 110 (COSAN)', 'Canal 1 (BRIDA)'),
            ('1', 'Canal 1 (BRIDA)', 'Canal 110 (COSAN)'),
            ('ambos', 'Canal 110 (COSAN)', None),
        ):
            params = {
                'inventario': self.inventario.pk,
                'canal': canal,
                'filtro': 'divergencias',
                'q': '00',
                'page': '9',
            }
            tela = self.client.get(reverse('confronto'), params)
            resposta = self._exportar_confronto(**params)
            self.assertEqual(resposta.status_code, 200, canal)
            self.assertIn(
                f'Confronto_Inventario_{self.inventario.pk}_',
                resposta['Content-Disposition'],
            )
            self.assertIn(
                timezone.localdate().strftime('%d-%m-%Y'),
                resposta['Content-Disposition'],
            )

            planilha = _planilha(resposta)
            cabecalhos = _cabecalhos(planilha)
            self.assertIn(coluna_presente, cabecalhos)
            if canal == 'ambos':
                self.assertIn('Canal 1 (BRIDA)', cabecalhos)
            else:
                self.assertNotIn(coluna_ausente, cabecalhos)
            self.assertEqual(planilha.freeze_panes, 'A2')
            self.assertTrue(planilha.auto_filter.ref.startswith('A1:'))
            self.assertGreaterEqual(planilha.column_dimensions['A'].width, 12)

            linhas_tela = list(tela.context['resultado'].linhas)
            self.assertGreaterEqual(len(linhas_tela), 2)
            self.assertEqual(len(_linhas(planilha)), len(linhas_tela))
            self.assertTrue(all(linha.possui_divergencia for linha in linhas_tela))

            for indice, linha in enumerate(linhas_tela, start=2):
                codigo = planilha.cell(indice, 1)
                self.assertEqual(codigo.value, linha.codigo_produto)
                self.assertEqual(codigo.number_format, FORMATO_TEXTO)
                self.assertEqual(codigo.data_type, 's')
                self.assertEqual(planilha.cell(indice, 2).value, linha.descricao)
                self.assertEqual(
                    planilha.cell(indice, 3).value,
                    linha.embalagem or '—',
                )
                self.assertEqual(
                    planilha.cell(indice, 4).value,
                    linha.setor or '—',
                )
                self._assert_numero(planilha, indice, cabecalhos, 'Físico', linha.fisico)
                if 'Canal 110 (COSAN)' in cabecalhos:
                    self._assert_numero(
                        planilha, indice, cabecalhos, 'Canal 110 (COSAN)', linha.canal_110,
                    )
                if 'Canal 1 (BRIDA)' in cabecalhos:
                    self._assert_numero(
                        planilha, indice, cabecalhos, 'Canal 1 (BRIDA)', linha.canal_1,
                    )
                self._assert_numero(
                    planilha, indice, cabecalhos, 'Total Contábil', linha.total_contabil,
                )
                self._assert_numero(
                    planilha, indice, cabecalhos, 'Diferença', linha.diferenca,
                )
                status = planilha.cell(indice, cabecalhos.index('Status') + 1)
                self.assertEqual(status.value, linha.status_label)

        self.assertContains(
            self.client.get(reverse('confronto'), {'inventario': self.inventario.pk, 'canal': '110'}),
            'Exportar Excel',
        )
        self.assertContains(
            self.client.get(reverse('confronto'), {'inventario': self.inventario.pk, 'canal': '1'}),
            'canal=1',
        )

        somente_codigo = self._exportar_confronto(canal='ambos', q='00055', filtro='todos')
        codigos = [linha[0] for linha in _linhas(_planilha(somente_codigo))]
        self.assertEqual(codigos, ['00055'])
        self.assertEqual(codigos[0], '00055')
        self.assertEqual(_planilha(somente_codigo).cell(2, 1).data_type, 's')

        self.assertEqual(InventarioItem.objects.count(), antes_itens)
        self.assertEqual(
            list(EstoqueSAP.objects.values_list('pk', 'canal_110', 'canal_1', 'total')),
            antes_sap,
        )

    def _assert_numero(self, planilha, linha, cabecalhos, titulo, esperado):
        cell = planilha.cell(linha, cabecalhos.index(titulo) + 1)
        self.assertEqual(cell.data_type, 'n', titulo)
        self.assertEqual(cell.number_format, FORMATO_NUMERO, titulo)
        self.assertEqual(_decimal(cell.value), esperado, titulo)

    def test_confronto_sem_inventario_nao_gera_arquivo(self):
        resposta = self.client.get(reverse('confronto_exportar'))
        self.assertRedirects(resposta, reverse('confronto'))

    def test_permissoes_da_exportacao(self):
        consulta, _ = criar_usuario_teste(
            username='consulta.excel',
            perfil=Usuario.Perfil.CONSULTA,
        )
        self.client.force_login(consulta)
        self.assertEqual(self._exportar_contagem().status_code, 200)
        self.assertEqual(self._exportar_confronto(canal='110').status_code, 200)

        operador, _ = criar_usuario_teste(
            username='operador.excel',
            perfil=Usuario.Perfil.OPERADOR,
        )
        self.client.force_login(operador)
        destino_pocket = reverse('pocket:selecionar')
        self.assertRedirects(
            self._exportar_contagem(),
            destino_pocket,
            fetch_redirect_response=False,
        )
        self.assertRedirects(
            self._exportar_confronto(),
            destino_pocket,
            fetch_redirect_response=False,
        )

        self.client.logout()
        self.assertEqual(self._exportar_contagem().status_code, 302)
        self.assertIn('/accounts/login/', self._exportar_contagem().url)
        self.assertEqual(self._exportar_confronto().status_code, 302)
        self.assertIn('/accounts/login/', self._exportar_confronto().url)
